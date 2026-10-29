"""Build 20 concise presentation slides + 4 backup slides from verified experiment data.

Run from any directory: python presentation/create_presentation.py
Text, diagrams, tables and charts remain editable in PowerPoint.
"""
from pathlib import Path
import json
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from examples.toy_bagging import build_toy_example


def read(path):
    return json.loads((ROOT / path).read_text(encoding="utf-8"))


R = read("outputs/results.json")
E = read("outputs/extensions/analysis.json")
C = read("outputs/extensions/controlled_comparison.json")
G = read("report/group_info.json")
T = build_toy_example()
if (R.get("seed") != 42 or R.get("n_estimators") != 51
        or R.get("selected_depth") != 3 or R.get("train_size") != 426
        or R.get("test_size") != 143
        or R.get("repeated_evaluation", {}).get("seeds") != list(range(42, 52))
        or [(m["tp"], m["fn"], m["fp"], m["tn"]) for m in R["test_metrics"]]
        != [(48, 5, 3, 87), (49, 4, 3, 87), (49, 4, 3, 87)]
        or C.get("protocol", {}).get("n_train") != 426
        or C.get("protocol", {}).get("n_test_excluded") != 143
        or any(r["n_estimators"] != 51 for r in C["summary"])
        or [r["n_estimators"] for r in E["tree_count_survey"]["summary"]]
        != [1, 5, 11, 21, 51, 101]):
    raise SystemExit("Cần kết quả mặc định: chạy de_5_ensemble_learning.py, "
                     "experiments/report_extensions.py và experiments/controlled_comparison.py trước.")
if not 1 <= len(G["members"]) <= 3:
    raise SystemExit("Kiểm tra thông tin nhóm trong report/group_info.json.")

P = Presentation()
P.slide_width, P.slide_height = Inches(13.333333), Inches(7.5)
P.core_properties.title = "Đề 5 — Ensemble Learning | Bagging cây quyết định"
P.core_properties.subject = "Báo cáo giữa kỳ — Nhập môn học máy"
P.core_properties.author = "; ".join(m["name"] for m in G["members"])
P.core_properties.keywords = "Bagging, bootstrap, hard voting, OOB, đối chứng"
NAVY, TEAL, BLUE = "142D46", "087F82", "416DC4"
CORAL, WHITE, BG = "C86449", "FFFFFF", "F4F7FA"
INK, GRAY, LINE, MINT = "192F43", "5B6B7B", "DCE4EB", "8BE0CF"
PALE, SOFT_RED = "E9F4F1", "FAEDE8"
FONT, MONO = "Noto Sans", "DejaVu Sans Mono"
NOTES, TITLES = [], []
MAIN_SLIDES, TOTAL_SLIDES, HANDOFF = 20, 24, 11


def rgb(s):
    return RGBColor.from_string(s)


def num(v, digits=2):
    return f"{v:.{digits}f}".replace(".", ",")


def pct(v):
    return num(100 * v) + "%"


def avg(row, key="balanced_accuracy"):
    return f"{num(100 * row[key + '_mean'])} ± {num(100 * row[key + '_std'])}"


def shape(s, x, y, w, h, fill=WHITE, stroke=None, kind=MSO_SHAPE.RECTANGLE):
    obj = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid(); obj.fill.fore_color.rgb = rgb(fill)
    if stroke:
        obj.line.color.rgb = rgb(stroke); obj.line.width = Pt(.8)
    else:
        obj.line.fill.background()
    # Remove theme effects rather than inheriting the template's drop shadows.
    obj._element.spPr.append(OxmlElement("a:effectLst"))
    style = obj._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None: effect.set("idx", "0")
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        obj.adjustments[0] = .08
    return obj


def txt(s, x, y, w, h, value, size=22, color=INK, bold=False,
        align=PP_ALIGN.LEFT, font=FONT, leading=1.12):
    obj = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = obj.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, text in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text; p.alignment = align
        p.font.name = font; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = rgb(color)
        p.space_before = Pt(0); p.space_after = Pt(0)
        p.line_spacing = leading
    return obj


def rule(s, x1, y1, x2, y2, color=LINE, width=1):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(width)
    ln._element.spPr.append(OxmlElement("a:effectLst"))
    style = ln._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None: effect.set("idx", "0")
    return ln


def arrow(s, x, y, width=.34, color=TEAL):
    return shape(s, x, y, width, .2, color, kind=MSO_SHAPE.RIGHT_ARROW)


def panel(s, x, y, w, h, fill=BG, stroke=None):
    return shape(s, x, y, w, h, fill, stroke, MSO_SHAPE.ROUNDED_RECTANGLE)


def slide(title, section, source="", dark=False, backup=False):
    s = P.slides.add_slide(P.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY if dark else WHITE)
    TITLES.append(title)
    txt(s, .62, .29, 10.8, .25, section.upper(), 10.5, MINT if dark else TEAL, True)
    txt(s, .62, .87, 12.08, .95, title, 30, WHITE if dark else NAVY, True)
    rule(s, .62, 7.02, 12.7, 7.02, "35516A" if dark else LINE, .7)
    txt(s, .62, 7.14, 3.75, .17, "ĐỀ 5 / ENSEMBLE LEARNING", 8.2, MINT if dark else GRAY)
    txt(s, 4.2, 7.14, 7.72, .17, source, 8.2, "B5CAD8" if dark else GRAY, align=PP_ALIGN.RIGHT)
    txt(s, 12.14, 7.1, .56, .25, f"{len(P.slides):02d}", 11, MINT if dark else TEAL, True, PP_ALIGN.RIGHT)
    # Quiet chapter marker, including a distinct appendix indicator.
    for i in range(4):
        shape(s, 11.67 + .27 * i, .36, .18, .035,
              (MINT if dark else TEAL) if backup or i == min((len(P.slides) - 1) // 5, 3)
              else ("35516A" if dark else LINE))
    # Standard OOXML Fade: manual advance, no automatic slide timing or sound.
    transition = OxmlElement("p:transition")
    transition.set("spd", "fast")
    transition.set("advClick", "1")
    transition.append(OxmlElement("p:fade"))
    s._element.insert_element_before(transition, "p:timing", "p:extLst")
    return s


def note(s, value):
    owner = "Nguyễn Nghĩa Hiếu" if len(P.slides) <= HANDOFF else "Trần Quốc Bảo"
    if len(P.slides) > MAIN_SLIDES: owner = "Cả hai — dự phòng khi hỏi đáp"
    value = f"Người nói gợi ý: {owner}.\n\n" + value
    s.notes_slide.notes_text_frame.text = value
    NOTES.append((len(P.slides), TITLES[-1], value))


def takeaway(s, value, dark=False, y=6.31):
    panel(s, .62, y, 12.08, .48, "213E57" if dark else PALE)
    txt(s, .82, y + .105, 11.65, .29, value, 15.5, MINT if dark else TEAL, True)


def label(s, x, y, value, color=TEAL, w=4):
    txt(s, x, y, w, .28, value.upper(), 11, color, True)


def table(s, x, y, widths, rows, row_h=.7, size=19, highlight=()):
    for ri, values in enumerate(rows):
        xx = x
        for ci, (val, w) in enumerate(zip(values, widths)):
            color = NAVY if ri == 0 else (PALE if ri in highlight else (WHITE if ri % 2 else BG))
            shape(s, xx, y + ri * row_h, w, row_h, color)
            txt(s, xx + .17, y + ri * row_h + .16, w - .34, row_h - .18, val,
                size if ri else size - 1, WHITE if ri == 0 else INK, ri == 0 or ri in highlight,
                PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER)
            xx += w
        if ri > 0: rule(s, x, y + (ri + 1) * row_h, x + sum(widths), y + (ri + 1) * row_h, LINE, .55)


def stat(s, x, y, value, title, detail="", color=TEAL, w=3.7):
    txt(s, x, y, w, .85, value, 45, color, True)
    txt(s, x, y + .95, w, .5, title, 19, INK, True)
    if detail: txt(s, x, y + 1.6, w, .8, detail, 16, GRAY)


def code(s, x, y, w, value, h=4.45):
    panel(s, x, y, w, h, NAVY)
    label(s, x + .28, y + .22, "PYTHON · TRÍCH LƯỢC", MINT)
    txt(s, x + .28, y + .79, w - .56, h - .93, value, 17.3, WHITE, font=MONO, leading=1.22)


def count_chart(s, summary):
    data = CategoryChartData()
    data.categories = [str(row["n_estimators"]) for row in summary]
    data.add_series("Balanced accuracy CV", [100 * row["balanced_accuracy_mean"] for row in summary])
    chart = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(.7), Inches(2.05), Inches(7.5), Inches(3.6), data).chart
    chart.has_legend = False
    chart.has_title = False
    chart.chart_style = 2
    chart.value_axis.minimum_scale = 90; chart.value_axis.maximum_scale = 98
    chart.value_axis.major_unit = 2
    chart.value_axis.tick_labels.font.name = FONT; chart.value_axis.tick_labels.font.size = Pt(13)
    chart.category_axis.tick_labels.font.name = FONT; chart.category_axis.tick_labels.font.size = Pt(14)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = rgb(LINE)
    series = chart.series[0]
    series.format.line.color.rgb = rgb(TEAL); series.format.line.width = Pt(2.3)
    series.marker.size = 8
    series.marker.style = XL_MARKER_STYLE.CIRCLE
    series.marker.format.fill.solid(); series.marker.format.fill.fore_color.rgb = rgb(TEAL)
    series.marker.format.line.color.rgb = rgb(WHITE)
    series.has_data_labels = True
    labels = series.data_labels
    labels.position = XL_LABEL_POSITION.ABOVE; labels.number_format = '0.00'
    labels.show_value = True
    labels.font.name = FONT; labels.font.size = Pt(14); labels.font.color.rgb = rgb(TEAL)
    for point, row in zip(series.points, summary):
        p = point.data_label.text_frame.paragraphs[0]
        p.text = num(100 * row["balanced_accuracy_mean"])
        p.font.name = FONT; p.font.size = Pt(14); p.font.color.rgb = rgb(TEAL)
    return chart


# 01 — Institutional cover with a restrained, editable ensemble diagram.
s = slide("", "Báo cáo giữa kỳ · Nhập môn học máy")
s.shapes.add_picture(str(ROOT / "presentation/assets/tdtu_logo.png"), Inches(.63), Inches(.86), width=Inches(1.3))
txt(s, 2.2, .98, 8.8, .65, "TRƯỜNG ĐẠI HỌC TÔN ĐỨC THẮNG\nKHOA CÔNG NGHỆ THÔNG TIN", 12, GRAY, True)
txt(s, .62, 2.0, 7.2, 1.7, "ENSEMBLE\nLEARNING", 48, NAVY, True, leading=1.0)
txt(s, .67, 4.02, 7.15, .9, "Bagging cây quyết định\ntrong phân loại u vú", 25, TEAL)
identities = [f"{m['name']}  ·  {m['student_id']}" for m in G["members"]]
identities.append("Giảng viên: " + G["lecturer"])
txt(s, .67, 5.61, 7.5, .95, "\n".join(identities), 13.5, GRAY, leading=1.35)
panel(s, 8.32, 2.02, 4.38, 4.66, NAVY)
label(s, 8.65, 2.4, "BOOTSTRAP → VOTING", MINT)
for yy, lab in [(3.02, "Cây 1"), (3.96, "Cây 2"), (4.90, "Cây B")]:
    panel(s, 8.67, yy, 1.45, .64, "27465F")
    txt(s, 8.82, yy + .16, 1.15, .32, lab, 17, WHITE)
    rule(s, 10.14, yy + .32, 11.08, 4.28, "5CAAA7", 1.5)
shape(s, 10.82, 3.79, 1.32, 1.0, TEAL, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 10.91, 4.05, 1.14, .45, "H(x)", 25, WHITE, True, PP_ALIGN.CENTER)
txt(s, 8.67, 6.0, 3.55, .35, "Nhiều cây · Một dự đoán", 15, MINT)
note(s, "Mở đầu bằng câu hỏi: khi một cây quyết định nhạy với dữ liệu, kết hợp nhiều cây có giúp cải thiện dự đoán không? Ba câu hỏi chính là: Bagging có cải thiện so với cây đơn; tăng số cây có tiếp tục có lợi; và đa dạng hơn có đủ giúp tổ hợp tốt hơn? Báo cáo tập trung vào Bagging. Nhóm tự triển khai bootstrap, hard voting và OOB; cây cơ sở dùng scikit-learn. Sơ đồ trên bìa minh họa cơ chế kết hợp, không biểu diễn những cây đã học cụ thể. Bộ slide gồm 20 trang chính và 4 trang dự phòng cho hỏi đáp. Thời lượng luyện tập gợi ý là 15–18 phút, không phải quy định của đề bài.")

# 02 — Landscape; the scope is explicit.
s = slide("Bagging trong bức tranh học tổ hợp", "01 / Cơ sở lý thuyết", "[1], [2], [3], [6], [7]")
for i, (title, desc, end) in enumerate([
    ("Bagging", "Bootstrap\nHọc độc lập", "TRỌNG TÂM TRIỂN KHAI"),
    ("Boosting", "Học tuần tự\nGiảm sai số / loss", "AdaBoost / Gradient Boosting"),
    ("Voting", "Kết hợp nhãn\nhoặc xác suất", "Hard / Soft"),
    ("Stacking", "Dự đoán OOF\n→ học tầng hai", "Meta-model"),
]):
    x = .62 + i * 3.08
    panel(s, x, 2.14, 2.84, 3.67, PALE if i == 0 else BG)
    shape(s, x + .22, 2.48, .38, .05, TEAL if i == 0 else BLUE)
    txt(s, x + .22, 2.83, 2.4, .5, title, 26, NAVY, True)
    txt(s, x + .22, 3.68, 2.42, 1.4, desc, 20, GRAY, leading=1.3)
    txt(s, x + .22, 5.28, 2.42, .4, end, 10.7, TEAL, True)
takeaway(s, "Trọng tâm: Bagging cây quyết định và quy tắc hard voting.")
note(s, "Bagging tạo nhiều mẫu bootstrap để học các mô hình cơ sở. Boosting xây dựng một chuỗi mô hình nhằm xử lý sai sót hoặc giảm hàm mất mát. Voting là quy tắc kết hợp nhãn hay xác suất; Stacking học một bộ kết hợp từ dự đoán ngoài fold. Các hướng này không hoàn toàn tách rời: Bagging cũng cần quy tắc tổng hợp. Bài chỉ thực nghiệm sâu Bagging và Random Forest. Boosting, Stacking được giới thiệu tổng quan; chi tiết minh họa nằm ở slide dự phòng 23.")

# 03 — Bagging flow.
s = slide("Bagging = Bootstrap + Aggregating", "02 / Cơ chế", "[1], [3] · ensemble_model.py")
for i, (title, body_) in enumerate([
    ("Giữ riêng test", "Chỉ lấy mẫu\ntừ tập train."),
    ("Tạo bootstrap", "N lượt lấy\ncó hoàn lại / cây."),
    ("Học B cây", "Một cây học\ntrên mỗi bootstrap."),
    ("Đếm phiếu", "Cùng một mẫu mới\n→ chọn lớp đa số."),
]):
    x = .62 + i * 3.1
    label(s, x, 2.38, f"BƯỚC {i + 1}", w=2.65)
    rule(s, x, 2.94, x + 2.65, 2.94, LINE, 2)
    txt(s, x, 3.28, 2.72, .65, title, 24, NAVY, True)
    txt(s, x, 4.27, 2.72, 1.15, body_, 21, GRAY)
    if i < 3: arrow(s, x + 2.72, 3.45, .25)
takeaway(s, "Mỗi bootstrap lấy N lượt có hoàn lại từ cùng tập train.")
note(s, "Mỗi bootstrap có N lượt lấy có hoàn lại, nên một quan sát có thể lặp hoặc không xuất hiện. Mỗi cây học trên bootstrap riêng rồi cùng dự đoán một quan sát mới. Các cây học độc lập về quy trình, không có nghĩa dự đoán của chúng độc lập thống kê. Mã minh họa huấn luyện tuần tự. Mô hình tham chiếu có N = 426, B = 51 và max_depth = 3. Số cây lẻ tránh hòa trong bài toán hai lớp; khi có nhiều lớp vẫn có thể hòa, mã chọn lớp đầu trong classes_.")

# 04 — One connected toy example, starting with the data.
s = slide("Cùng 5 quan sát, tạo 3 bootstrap khác nhau", "02 / Ví dụ tính tay", "examples/toy_bagging.py · Dữ liệu giả lập")
label(s, .67, 1.98, "DỮ LIỆU GỐC")
for i, (x, y) in enumerate(zip(T["dataset"]["feature_values"], T["dataset"]["labels"])):
    xx = 3.0 + i * 1.87
    txt(s, xx, 1.95, 1.6, .44, f"x = {int(x)}", 22, NAVY, True, PP_ALIGN.CENTER)
    txt(s, xx, 2.48, 1.6, .35, "Lành" if y else "Ác", 17, TEAL if y else CORAL, align=PP_ALIGN.CENTER)
rows = [["Cây", "Chỉ số bootstrap (từ 0)", "Quan sát vắng mặt"]]
for i, tr in enumerate(T["trees"]):
    rows.append([f"h{i + 1}", str(tr["bootstrap_indices"]), ", ".join(str(k) for k in tr["omitted_indices"])])
table(s, .67, 3.2, [1.6, 6.5, 3.85], rows, .65, 20)
takeaway(s, "Có lặp, có quan sát vắng mặt; tổng số lượt lấy vẫn bằng N.")
note(s, "Ví dụ gồm x = [1, 2, 3, 4, 5] và y = [1, 1, 0, 0, 0], với nhãn 1 là lành, nhãn 0 là ác. Các chỉ số hàng bắt đầu từ 0. Ba bootstrap được chọn cố định để dễ tính tay: h₁ vắng hàng 3, h₂ vắng hàng 2 và h₃ vắng hàng 0. Đây là dữ liệu giả lập, không đại diện cho một đặc trưng y tế cụ thể. Trong thí nghiệm chính, bootstrap vẫn được sinh ngẫu nhiên có hoàn lại theo seed.")

# 05 — Gini and a minimal decision tree.
s = slide("Cây h₁ chọn ngưỡng làm giảm Gini nhiều nhất", "02 / Ví dụ tính tay", "examples/toy_bagging.py · [5]")
label(s, .67, 2.1, "TRƯỚC KHI CHIA")
txt(s, .67, 2.67, 5.5, .52, "3 lành + 2 ác", 30, NAVY, True)
txt(s, .67, 3.44, 5.5, .9, "Gini = 1 − (3/5)² − (2/5)²\n        = 0,48", 24, TEAL, leading=1.3)
table(s, .67, 4.82, [2.55, 2.6], [["Ngưỡng", "Gini sau chia"], ["1,5 / 2,5 / 4,0", "0,2667 / 0 / 0,30"]], .58, 16)
panel(s, 6.8, 2.05, 5.88, 3.85, BG)
panel(s, 8.33, 2.48, 2.8, .7, NAVY)
txt(s, 8.55, 2.65, 2.36, .38, "x ≤ 2,5 ?", 24, WHITE, True, PP_ALIGN.CENTER)
rule(s, 9.0, 3.2, 8.1, 4.15, LINE, 2); rule(s, 10.45, 3.2, 11.36, 4.15, LINE, 2)
txt(s, 7.58, 3.57, 1.2, .3, "ĐÚNG", 11, TEAL, True)
txt(s, 10.7, 3.57, 1.2, .3, "SAI", 11, CORAL, True)
txt(s, 7.18, 4.26, 2.6, .48, "3 lành", 27, TEAL, True, PP_ALIGN.CENTER)
txt(s, 9.98, 4.26, 2.4, .48, "2 ác", 27, CORAL, True, PP_ALIGN.CENTER)
txt(s, 7.28, 5.05, 5.0, .4, "Cả hai nhánh: Gini = 0", 20, GRAY, align=PP_ALIGN.CENTER)
takeaway(s, "Ngưỡng 2,5: Gini giảm từ 0,48 xuống 0.")
note(s, "Bootstrap của h₁ có các giá trị [1, 1, 3, 5, 2], gồm ba nhãn lành và hai nhãn ác. Gini ban đầu bằng 0,48. Các ngưỡng ứng viên là 1,5; 2,5 và 4,0, cho Gini sau chia lần lượt 0,2667; 0 và 0,30. Ngưỡng 2,5 có mức giảm lớn nhất nên được chọn. h₂ thiếu x = 3 nên học ngưỡng 3,0; h₃ có ngưỡng 2,5. Phép tính được đối chiếu với cây một mức của scikit-learn, không phải bộ học cây do nhóm tự cài từ đầu.")

# 06 — Voting remains linked to exactly those three fitted trees.
s = slide("Mẫu x = 2,8 nhận hai phiếu ác và một phiếu lành", "02 / Ví dụ tính tay", "Cùng ba bootstrap ở slide 4 · examples/toy_bagging.py")
for i, tr in enumerate(T["trees"]):
    x = .67 + i * 4.13; pred = tr["query_prediction"]
    panel(s, x, 2.22, 3.79, 2.65, BG)
    label(s, x + .25, 2.51, f"CÂY h{i + 1}")
    rel = "≤" if 2.8 <= tr["threshold"] else ">"
    txt(s, x + .25, 3.12, 3.3, .47, f"2,8 {rel} {num(tr['threshold'], 1)}", 28, NAVY, True)
    txt(s, x + .25, 3.94, 3.3, .5, "Lành tính" if pred else "Ác tính", 27, TEAL if pred else CORAL, True)
txt(s, .7, 5.17, 4.3, .8, "2 : 1", 46, CORAL, True)
txt(s, 4.72, 5.34, 7.7, .6, "Hard voting chọn ác tính (0).", 25, NAVY, True)
takeaway(s, "Ví dụ giả lập, chưa có nhãn thật; tỷ lệ phiếu không phải xác suất bệnh.")
note(s, "Đưa cùng quan sát mới x = 2,8 qua ba cây. Với h₁ và h₃, giá trị này lớn hơn ngưỡng 2,5 nên dự đoán ác, nhãn 0. Với h₂, 2,8 nhỏ hơn hoặc bằng ngưỡng 3,0 nên dự đoán lành, nhãn 1. Tổ hợp có hai phiếu ác và một phiếu lành, vì vậy chọn nhãn 0. Đây chỉ là minh họa quy tắc đa số. Mẫu mới chưa có nhãn thật, nên chưa thể kết luận dự đoán đúng hay sai; tỷ lệ phiếu cũng chưa phải xác suất bệnh.")

# 07 — Rules of aggregation.
s = slide("Hard và soft voting có thể chọn khác lớp", "02 / Quy tắc tổng hợp", "Ví dụ giả định · [3]")
table(s, .67, 2.15, [4.45, 2.5, 2.5, 2.5], [
    ["Đầu ra từng mô hình", "Cây 1", "Cây 2", "Cây 3"],
    ["P(ác tính)", "0,51", "0,51", "0,01"],
    ["Nhãn dự đoán", "Ác", "Ác", "Lành"],
], .67, 20)
for x, title, desc, col in [(.67, "HARD → ÁC", "2 phiếu ác / 3 cây", CORAL),
                             (6.85, "SOFT → LÀNH", "P(ác) trung bình = 0,343", TEAL)]:
    panel(s, x, 4.7, 5.77, 1.28, BG)
    txt(s, x + .23, 4.91, 5.3, .43, title, 24, col, True)
    txt(s, x + .23, 5.48, 5.3, .32, desc, 18, GRAY)
takeaway(s, "Hard đếm nhãn; soft trung bình xác suất.")
note(s, "Đây là ví dụ giả định để phân biệt hai quy tắc tổng hợp, khác với ví dụ tính tay trước. Hard voting chỉ đếm nhãn nên chọn ác với hai trong ba phiếu. Soft voting lấy trung bình xác suất: (0,51 + 0,51 + 0,01)/3 ≈ 0,3433 cho lớp ác, tương ứng khoảng 0,6567 cho lớp lành, nên chọn lành. Hai quy tắc có thể cho kết quả khác; soft không tự động tốt hơn hard. Phần đối chứng ở slide 17 kiểm tra chúng trên cùng các cây đã học thực tế.")

# 08 — OOB and the finite bootstrap interpretation.
s = slide("OOB chỉ dùng phiếu của cây chưa học quan sát đó", "02 / Đánh giá bổ sung", "outputs/results.json → oob · [1]")
panel(s, .67, 2.08, 6.06, 3.88, BG)
for j, (lab, has_seen) in enumerate([("Cây 1: có mẫu i", True), ("Cây 2: vắng mẫu i", False), ("Cây 3: vắng mẫu i", False)]):
    yy = 2.48 + j * 1.0
    txt(s, .96, yy, 3.6, .43, lab, 21, INK)
    txt(s, 4.56, yy, 1.8, .43, "Không tính" if has_seen else "Được bỏ phiếu", 16.5, GRAY if has_seen else TEAL, True)
stat(s, 7.17, 2.2, pct(R["oob"]["accuracy"]), "Accuracy OOB", "401/426 dự đoán đúng", w=5)
txt(s, 7.17, 5.0, 5.1, .8, "Coverage 100%\nMỗi mẫu: 8–34 phiếu OOB", 21, NAVY, True, leading=1.3)
takeaway(s, "OOB là đánh giá bổ sung, được tính sau chọn độ sâu.")
note(s, "Mỗi quan sát chỉ nhận phiếu OOB từ các cây chưa lấy nó vào bootstrap. Hàng chưa có phiếu được giữ không xác định, tránh dùng argmax trên hàng toàn 0. Lần tham chiếu có coverage 100% và 401/426 dự đoán đúng. Với N = 426, xác suất vắng trong một bootstrap khoảng 36,74%. Số quan sát khác nhau kỳ vọng là 269,47, thực tế trung bình 268,04; số phiếu OOB kỳ vọng mỗi mẫu là 18,74, thực tế 18,91. Coverage khác accuracy. OOB được tính sau chọn độ sâu nên chưa độc lập hoàn toàn với tuning.")

# 09 — Data and train-only model selection.
s = slide("569 mẫu: học trên train, đánh giá trên test", "03 / Dữ liệu và giao thức", "Breast Cancer Wisconsin Diagnostic · seed 42 · [4]")
for x, value, caption in [(.73, "569", "quan sát"), (4.83, "30", "đặc trưng số"), (8.93, "2", "lớp: ác (0), lành (1)")]:
    txt(s, x, 2.03, 3.4, .85, value, 44, TEAL, True)
    txt(s, x, 3.04, 3.4, .4, caption, 19, GRAY)
panel(s, .67, 3.87, 8.1, 2.13, PALE)
label(s, .94, 4.13, "TRAIN 426 MẪU · 75%", TEAL, 7.4)
txt(s, .95, 4.86, 3.0, .5, "CV 5 fold", 27, NAVY, True)
arrow(s, 4.2, 5.04, .4)
txt(s, 5.02, 4.86, 3.1, .5, "Độ sâu 3", 27, TEAL, True)
panel(s, 9.02, 3.87, 3.64, 2.13, NAVY)
label(s, 9.3, 4.13, "TEST GIỮ RIÊNG", MINT, 3.05)
txt(s, 9.3, 4.86, 3.05, .65, "143 mẫu", 30, WHITE, True)
takeaway(s, "Chọn độ sâu bằng cây đơn trên train, rồi dùng chung cho ba mô hình.")
note(s, "Dữ liệu gồm 569 quan sát, 212 mẫu ác và 357 mẫu lành, với 30 đặc trưng số đã trích xuất từ nhân tế bào. Ác tính có nhãn 0 và được quy ước là lớp dương. Seed 42 chia có phân tầng theo tỷ lệ 75/25: train 426 gồm 159 ác, 267 lành; test 143 gồm 53 ác, 90 lành. Chỉ trên train, CV phân tầng năm fold thử độ sâu 2–6; độ sâu 3 đạt balanced accuracy CV cao nhất, 93,73%. Độ sâu được chọn bằng cây đơn rồi dùng chung cho ba mô hình; chưa tối ưu riêng từng ensemble. Test không tham gia lựa chọn. Mỗi seed 42–51 lặp lại quy trình này.")

# 10 — Readable source code, one operation per slide.
s = slide("Bootstrap được triển khai trực tiếp trong Python", "04 / Triển khai", "Trích lược ensemble_model.py · ManualBagging.fit")
code(s, .67, 2.0, 8.21, '''indices = rng.integers(
    0, n_samples, size=n_samples
)
tree = DecisionTreeClassifier(
    max_depth=self.max_depth,
    random_state=int(rng.integers(0, 2**31 - 1)),
)
tree.fit(X[indices], encoded[indices])
self.estimators_.append(tree)''', 4.11)
for i, (a, b) in enumerate([("Lấy có hoàn lại", "Có thể lặp chỉ số."), ("Seed riêng từng cây", "Hỗ trợ tái lập."), ("Đúng cặp X và y", "Dùng cùng indices.")]):
    yy = 2.28 + 1.17 * i
    label(s, 9.25, yy, f"0{i+1}", TEAL, 3.1)
    txt(s, 9.25, yy + .32, 3.15, .42, a, 20, NAVY, True)
    txt(s, 9.25, yy + .85, 3.15, .35, b, 15.5, GRAY)
takeaway(s, "Tự triển khai Bagging; cây cơ sở dùng scikit-learn.")
note(s, "Đoạn trích nằm trong phương thức fit; X ở đây chỉ chứa train. Nhãn encoded được tạo bằng np.unique(y, return_inverse=True), giúp đưa các lớp về chỉ số liên tiếp. Bộ sinh rng khởi tạo từ random_state, rồi sinh n_samples chỉ số có hoàn lại và seed cho từng cây. Mỗi cây học X và nhãn trên cùng chỉ số bootstrap. Mã đầy đủ còn lưu estimators_samples_ để truy vết và tính OOB. Cây DecisionTreeClassifier dùng thư viện; phần nhóm tự triển khai là tổ chức lấy mẫu, nhiều cây và quản lý phiếu.")

# 11 — Voting, with the class mapping retained.
s = slide("Hard voting đếm phiếu rồi ánh xạ về nhãn gốc", "04 / Triển khai", "Trích lược ensemble_model.py · votes / predict")
code(s, .67, 2.0, 8.21, '''counts = np.zeros(
    (len(X), len(self.classes_)), dtype=np.int64
)
rows = np.arange(len(X))
for tree in self.estimators_:
    predicted = tree.predict(X).astype(int)
    np.add.at(counts, (rows, predicted), 1)
labels = self.classes_[counts.argmax(axis=1)]''', 4.11)
txt(s, 9.25, 2.36, 3.12, .85, "Mỗi hàng =\nmột quan sát", 23, NAVY, True)
txt(s, 9.25, 3.72, 3.12, .85, "Mỗi cột =\nmột lớp", 23, NAVY, True)
txt(s, 9.25, 5.06, 3.12, .65, "Tổng phiếu / hàng = B", 18, TEAL, True)
takeaway(s, "Mỗi cây bỏ một phiếu; argmax chọn lớp có nhiều phiếu nhất.")
note(s, "Các cây học nhãn mã hóa nên predict trả về chỉ số cột lớp. np.add.at cộng một phiếu vào đúng hàng và cột; argmax chọn lớp nhiều phiếu nhất, rồi self.classes_ ánh xạ về nhãn gốc. Trong mã đầy đủ, votes và predict là hai phương thức riêng; slide gộp để giải thích. Bộ mã đạt 26/26 kiểm thử, gồm hard/soft khác nhãn, cây học một lớp và phép đo bất đồng. Chuyển lời cho thành viên thứ hai: từ cơ chế đã kiểm chứng, phần tiếp theo xem các cây hoạt động ra sao trên dữ liệu thực.")

# 12 — Holdout scores and error counts on one slide.
s = slide("Seed 42: Bagging đúng thêm 1/143 mẫu test", "05 / Kết quả gốc", "outputs/results.json · ác tính (0) là lớp dương")
rows = [["Mô hình", "Accuracy", "Recall ác"]]
for name, m in zip(["Cây đơn", "Bagging", "Random Forest"], R["test_metrics"]):
    rows.append([name, pct(m["accuracy"]), pct(m["recall_malignant"])])
table(s, .67, 2.61, [2.8, 1.65, 1.65], rows, .66, 18.5, (2,))
label(s, 7.06, 2.03, "MA TRẬN BAGGING", TEAL, 5.4)
txt(s, 8.42, 2.61, 3.2, .3, "DỰ ĐOÁN", 11, GRAY, True, PP_ALIGN.CENTER)
for x, name in [(8.46, "Ác"), (10.07, "Lành")]:
    txt(s, x, 3.0, 1.5, .35, name, 18, INK, align=PP_ALIGN.CENTER)
for r, row in enumerate([[49, 4], [3, 87]]):
    txt(s, 7.0, 3.72 + r * .98, 1.25, .4, "Thật ác" if r == 0 else "Thật lành", 16.5, GRAY)
    for c, value in enumerate(row):
        xx, yy = 8.43 + 1.61 * c, 3.45 + .98 * r
        shape(s, xx, yy, 1.53, .92, PALE if r == c else SOFT_RED)
        txt(s, xx, yy + .16, 1.53, .55, str(value), 31, TEAL if r == c else CORAL, True, PP_ALIGN.CENTER)
txt(s, .76, 5.65, 5.9, .43, "Bỏ sót ác: 5 → 4", 25, TEAL, True)
txt(s, 7.08, 5.65, 5.25, .43, "Dương tính giả: 3", 21, GRAY)
takeaway(s, "Bagging vẫn bỏ sót 4/53 mẫu ác; một lần chia chưa đủ kết luận ưu thế tổng quát.")
note(s, "Hàng của ma trận là nhãn thật, cột là dự đoán. Bagging có TP = 49, FN = 4, FP = 3 và TN = 87; cây đơn có TP = 48, FN = 5, FP = 3 và TN = 87. Bagging đúng 136/143, cây đơn đúng 135/143, chênh khoảng 0,70 điểm phần trăm accuracy. Precision ác của Bagging là 94,23%, F1 ác là 93,33%; xem phụ lục chỉ số khi được hỏi. Random Forest có cùng chỉ số tổng hợp ở seed 42 nhưng sai trên những quan sát khác: hai mô hình khác dự đoán ở hàng 291 và 128. Không suy ra tương đương từ điểm tổng hợp. ManualBagging dùng toàn bộ đặc trưng và hard voting; RF sklearn dùng ngẫu nhiên hóa đặc trưng và trung bình xác suất.")

# 13 — Ten paired holdouts.
s = slide("Bagging cao hơn cây đơn ở 9/10 lần chia", "05 / Mười lần chia gốc", "outputs/results.json · controlled_comparison.json · seed 42–51")
rows = [["Mô hình", "Accuracy (%)", "Recall ác (%)"]]
for name, m in zip(["Cây đơn", "Bagging", "Random Forest"], R["repeated_evaluation"]["summary"]):
    rows.append([name, avg(m, "accuracy"), avg(m, "recall_malignant")])
table(s, .67, 2.1, [4.45, 3.75, 3.75], rows, .72, 21, (2,))
txt(s, .75, 5.3, 4.0, .7, "+1,96 đpt", 35, TEAL, True)
txt(s, 5.0, 5.24, 7.45, .85, "Accuracy trung bình so với cây đơn.\nSo với RF: chưa có ưu thế nhất quán.", 20, GRAY, leading=1.3)
takeaway(s, "Trung bình ± SD mẫu; các test có thể giao nhau. SD chưa phải khoảng tin cậy.")
note(s, "Trung bình được trình bày theo phần trăm; SD mẫu theo điểm phần trăm. Mỗi seed đều chia lại 75/25 có phân tầng và chọn độ sâu trên train. Khi so accuracy trong cùng seed, Bagging cao hơn cây đơn 9 lần, bằng 1 lần, thấp hơn 0 lần; chênh lệch trung bình khoảng +1,96 điểm phần trăm. So với RF, các số tương ứng là 2, 4 và 4, chênh lệch trung bình −0,28 điểm phần trăm. Các holdout có thể giao nhau, nên không coi là mười phép thử độc lập hay suy ra giá trị p từ số lần thắng.")

# 14 — Editable chart, with precise scope and SD beside it.
s = slide("Tăng số cây có lợi, nhưng chưa tăng mãi", "06 / Khảo sát bổ sung trên train", "outputs/extensions/analysis.json · CV ngoài 5 fold · không dùng test")
survey = E["tree_count_survey"]["summary"]
count_chart(s, survey)
txt(s, .92, 1.94, 6.8, .28, "BALANCED ACCURACY CV (%)", 11, GRAY, True)
txt(s, 3.0, 5.75, 4.1, .3, "Số cây B (các cấu hình đã xét)", 13, GRAY, align=PP_ALIGN.CENTER)
panel(s, 8.77, 2.13, 3.88, 3.88, BG)
label(s, 9.05, 2.44, "51 CÂY", TEAL, 3)
txt(s, 9.05, 3.04, 3.3, .75, avg(survey[4]), 31, TEAL, True)
txt(s, 9.05, 4.0, 3.15, .6, "BA CV (%) · trung bình ± SD", 17, GRAY)
txt(s, 9.05, 4.84, 3.15, .85, "101 cây: 95,07%\nChưa cải thiện thêm.", 20, NAVY, True, leading=1.3)
takeaway(s, "Khảo sát trên train; 51 cây chưa được chứng minh tối ưu tổng quát.")
note(s, "Khảo sát chỉ dùng 426 mẫu train của seed 42, với năm fold ngoài. Trong phần học của mỗi fold, CV năm fold chọn độ sâu bằng cây đơn; các độ sâu là [2, 6, 2, 2, 6]. Các cấu hình B = [1, 5, 11, 21, 51, 101] dùng chung độ sâu và chuỗi seed tương ứng. Balanced accuracy trung bình tăng đến 95,26% ở 51 cây, rồi còn 95,07% ở 101 cây. SD tại 51 cây là 2,45 điểm phần trăm, không phải khoảng tin cậy. Đây là khảo sát bổ sung, chưa chứng minh 51 tối ưu tổng quát.")

# 15 — Error case as a balanced vote, not a confidence gauge.
s = slide("Tổ hợp 51 cây vẫn có thể dự đoán sai", "06 / Phân tích mẫu 291", "outputs/extensions/analysis.json → reference_error")
label(s, .7, 2.12, "NHÃN THẬT: LÀNH TÍNH", TEAL, 7)
txt(s, .7, 2.77, 7.8, .62, "27 phiếu ác / 24 phiếu lành", 32, NAVY, True)
shape(s, .72, 3.93, 11.88 * 27 / 51, .76, CORAL)
shape(s, .72 + 11.88 * 27 / 51, 3.93, 11.88 * 24 / 51, .76, TEAL)
txt(s, .87, 4.09, 5.9, .4, "27 · Ác", 22, WHITE, True)
txt(s, 7.05, 4.09, 5.3, .4, "24 · Lành", 22, WHITE, True, PP_ALIGN.RIGHT)
txt(s, .75, 5.17, 6.9, .65, "Đa số chọn ác → dương tính giả", 25, CORAL, True)
txt(s, 9.02, 5.12, 3.3, .9, "Biên phiếu\n3/51 ≈ 5,88%", 19, GRAY, align=PP_ALIGN.RIGHT)
takeaway(s, "Mẫu đầu tiên có hai phía phiếu, được chọn theo thứ tự test.")
note(s, "Quan sát được chọn có chỉ số test 2, tức mẫu thứ ba khi bắt đầu đếm từ 0, tương ứng hàng dữ liệu 291. Nhãn thật là lành tính, nhãn 1. Bagging nhận 27 phiếu ác và 24 phiếu lành, nên chọn nhãn 0 và tạo lỗi dương tính giả. Biên phiếu bằng 3/51 ≈ 0,0588; chênh lệch hai tỷ lệ là 5,88 điểm phần trăm, chưa phải xác suất bệnh đã hiệu chỉnh. RF dự đoán đúng mẫu này nhưng sai hàng 128, nên tổng FP vẫn bằng 3. Truy vết hai cây ở slide 21.")

# 16 — Feature ranges, each on its own labelled scale.
s = slide("Mẫu sai nằm giữa các vùng trung tâm của hai lớp", "06 / Giải thích hành vi mô hình", "Đặc trưng chọn theo tần suất nút gốc trên train · outputs/extensions/analysis.json")
for j, f in enumerate(E["reference_error"]["comparisons"]):
    yy = 2.1 + j * 1.29
    txt(s, .71, yy + .12, 3.75, .45, f["feature"], 19, NAVY, True)
    low, high = f["benign_q25"], f["malignant_q75"]
    def xx(v): return 5.2 + 5.6 * (v - low) / (high - low)
    rule(s, 5.0, yy + .43, 11.15, yy + .43, LINE, 1)
    for prefix, col in [("benign", TEAL), ("malignant", CORAL)]:
        q1, med, q3 = (f[prefix + k] for k in ("_q25", "_median", "_q75"))
        shape(s, xx(q1), yy + .26, xx(q3) - xx(q1), .34, PALE if prefix == "benign" else SOFT_RED)
        rule(s, xx(med), yy + .22, xx(med), yy + .64, col, 2)
    shape(s, xx(f["value"]) - .07, yy + .34, .14, .18, NAVY, kind=MSO_SHAPE.DIAMOND)
    txt(s, 11.36, yy + .11, 1.16, .42, num(f["value"], 5 if j == 1 else 2), 18, NAVY, True, PP_ALIGN.RIGHT)
    txt(s, 5.05, yy + .77, 1.2, .25, num(low, 3 if j == 1 else 1), 11, GRAY)
    txt(s, 10.0, yy + .77, 1.15, .25, num(high, 3 if j == 1 else 1), 11, GRAY, align=PP_ALIGN.RIGHT)
for x, col, lab in [(1.0, TEAL, "Lành: hộp Q1–Q3"), (4.65, CORAL, "Ác: hộp Q1–Q3"), (8.4, NAVY, "◆ Mẫu 291")]:
    shape(s, x, 6.08, .18, .12, col)
    txt(s, x + .3, 6.0, 3.6, .3, lab, 12, GRAY)
takeaway(s, "Mỗi hàng có thang riêng; ba đặc trưng chỉ cung cấp một phần giải thích.", y=6.48)
note(s, "Ba đặc trưng được chọn từ tần suất xuất hiện tại nút gốc trên train: worst perimeter ở 16 cây, worst concave points ở 14 cây và worst radius ở 9 cây. Giá trị của mẫu 291 lần lượt là 109,10; 0,14890 và 16,25, đều nằm giữa trung vị hai lớp, đồng thời giữa Q3 lớp lành và Q1 lớp ác. Hộp biểu diễn Q1–Q3, vạch là trung vị, hình thoi là mẫu lỗi. Mỗi hàng có thang đo riêng, không so chiều dài giữa các đặc trưng. Phân tích này chưa xác định nguyên nhân y khoa.")

# 17 — A controlled comparison: design and result together.
s = slide("Hard/soft gần nhau; sqrt chưa cải thiện", "06 / Đối chứng trên train", "controlled_comparison_summary.csv · BA CV trung bình ± SD qua 5 fold")
txt(s, .75, 2.03, 11.85, .5, "426 train · 51 cây · Cùng fold, bootstrap và seed từng cây", 20, GRAY)
rows = [["Xét đặc trưng", "Tổng hợp", "Balanced accuracy CV (%)"]]
for r in C["summary"]:
    rows.append(["Toàn bộ" if r["features"] == "all" else "sqrt", r["aggregation"].capitalize(), avg(r)])
table(s, .67, 2.79, [3.3, 2.75, 5.9], rows, .61, 20, (1, 2))
takeaway(s, "Nhóm sqrt tự xây khác RF sklearn; SD giữa các fold chưa phải khoảng tin cậy.")
note(s, "Đối chứng 2 × 2 tách hai yếu tố: xét toàn bộ đặc trưng hoặc sqrt; hard voting hoặc trung bình xác suất. Chỉ dùng 426 mẫu train của seed 42, không dùng test 143. Cùng năm fold ngoài, trong phần học mỗi fold, CV năm fold chọn độ sâu bằng cây đơn rồi dùng chung cho bốn cấu hình. Giữ nguyên bootstrap và seed từng cây giữa hai nhóm đặc trưng; hard/soft dùng chính cùng các cây đã học. All hard/soft đều đạt 95,26 ± 2,45%; sqrt hard 93,62 ± 3,48%, sqrt soft 93,44 ± 3,77%. All hard/soft trùng nhãn trên 426 lượt dự đoán ngoài fold; sqrt khác ở một quan sát. Giữ hard và đổi all sang sqrt, BA giảm khoảng 1,63 điểm phần trăm. Chưa có kiểm định ý nghĩa thống kê. Đây là phân tích bổ sung sau thí nghiệm gốc, chưa tối ưu riêng mỗi biến thể.")

# 18 — One mechanism insight rather than another dense result table.
s = slide("Đa dạng hơn chưa đủ: cây cơ sở cũng phải đủ tốt", "06 / Bài học từ đối chứng", "controlled_comparison_summary.csv · số trung bình qua 5 fold", dark=True)
for x, title, before, after, detail, col in [
    (.74, "BẤT ĐỒNG GIỮA CÁC CÂY", "7,99%", "10,52%", "D tăng khi dùng sqrt", MINT),
    (7.12, "CHẤT LƯỢNG MỘT CÂY", "91,67%", "89,75%", "Balanced accuracy trung bình giảm", "F0AB92"),
]:
    label(s, x, 2.18, title, col, 5.3)
    txt(s, x, 3.08, 5.3, .77, before, 43, WHITE, True)
    txt(s, x, 4.05, 5.3, .77, "→ " + after, 43, col, True)
    txt(s, x, 5.25, 5.3, .5, detail, 19, "C5D6E1")
takeaway(s, "Tổ hợp hard: 95,26% → 93,62%, trong cấu hình đang khảo sát.", True)
note(s, "D là tỷ lệ các cặp cây dự đoán khác nhãn, trung bình trên các quan sát và fold; chưa phải tương quan lỗi hay mức độc lập. Với hai lớp và 51 cây, D tối đa khoảng 50,98%. Ngẫu nhiên hóa đặc trưng làm tăng bất đồng nhưng giảm chất lượng cây cơ sở; balanced accuracy tổ hợp hard cũng giảm. Điều này chưa chứng minh xét toàn bộ đặc trưng luôn tốt hơn. Công thức phương sai σ²[ρ + (1 − ρ)/B] chỉ tạo trực giác khi các dự đoán số cùng phương sai và tương quan; không suy ρ từ D.")

# 19 — Three concise limitations and an accessible reproducibility anchor.
s = slide("Tái lập được; kết luận có giới hạn", "07 / Thảo luận", "Mã, cấu hình và kết quả được cung cấp trong repository")
for i, (title, desc) in enumerate([
    ("Chưa tuning riêng", "Độ sâu chung\ncho các mô hình."),
    ("Các holdout giao nhau", "SD mô tả dao động,\nchưa phải khoảng tin cậy."),
    ("Phân tích bổ sung", "Train seed 42;\nchưa có dữ liệu ngoài."),
]):
    x = .68 + i * 4.13
    panel(s, x, 2.15, 3.78, 3.12, BG)
    txt(s, x + .24, 2.52, 3.3, .8, title, 24, NAVY, True)
    txt(s, x + .24, 3.69, 3.3, 1.45, desc, 18, GRAY, leading=1.3)
txt(s, .75, 5.43, 3.3, .6, "26/26 kiểm thử", 28, TEAL, True)
obj = txt(s, 5.0, 5.55, 7.35, .45, "github.com/Huyngaongow/hocmaygk", 21, BLUE)
obj.click_action.hyperlink.address = "https://github.com/Huyngaongow/hocmaygk"
takeaway(s, "Mục đích học thuật; chưa đánh giá khả năng chẩn đoán lâm sàng.")
note(s, "Bộ mã có 26 kiểm thử giúp phát hiện lỗi cài đặt, không thay thế đánh giá ngoài mẫu. Gói nộp kèm chương trình, kết quả và báo cáo. Chạy thí nghiệm gốc để tạo results.json trước hai script bổ sung; dùng các phiên bản thư viện trong requirements.txt. Học tuần tự B cây có chi phí xấp xỉ B lần chi phí một cây; dự đoán K quan sát, độ sâu d có bậc O(KBd). Đây là phân tích cấu trúc, chưa đo thời gian thực. Demo nhanh lưu ở outputs/quick để giữ riêng kết quả mặc định.")

# 20 — Answer the three questions, then stop the main talk.
s = slide("Ba kết luận từ dữ liệu đã khảo sát", "07 / Kết luận và trao đổi", "Cảm ơn thầy và các bạn · Slide 21–24 dành cho hỏi đáp", dark=True)
for i, (headline, detail) in enumerate([
    ("Bagging cải thiện so với cây đơn", "9/10 lần chia cao hơn cây đơn; trung bình +1,96 đpt."),
    ("Số cây cần được kiểm tra", "51 → 101 cây: BA CV chưa tăng trong khảo sát."),
    ("Đa dạng đi cùng chất lượng", "Khác nhau hơn chưa chắc dự đoán tốt hơn."),
]):
    yy = 2.07 + i * 1.42
    txt(s, .75, yy, .6, .55, f"0{i+1}", 24, MINT, True)
    txt(s, 1.7, yy, 10.6, .52, headline, 27, WHITE, True)
    txt(s, 1.7, yy + .68, 10.6, .5, detail, 19, "C5D6E1")
note(s, "Kết thúc bằng ba ý trả lời câu hỏi nghiên cứu. Bagging liên hệ bootstrap với tổng hợp dự đoán; cải thiện so với cây đơn trong các lần chia đã khảo sát; nhưng tăng số cây hoặc mức bất đồng chưa tự động tạo mô hình tốt hơn. Chưa thể khẳng định Bagging vượt trội nhất quán so với RF, 51 cây tối ưu tổng quát hay mô hình đủ hiệu quả lâm sàng. Phần trình bày chính dừng tại đây. Hai thành viên cùng tham gia hỏi đáp và chỉ mở slide dự phòng phù hợp với câu hỏi.")

# 21 — Backup: two true paths with honest leaf counts.
s = slide("A. Cùng mẫu 291, hai cây đi tới hai nhãn", "Phụ lục / Truy vết cây", "outputs/extensions/analysis.json · representative_paths", backup=True)
for x, title, lines, final, col in [
    (.68, "CÂY 1 → ÁC", ["worst radius ≤ 16,79500", "worst concave points > 0,14315", "worst symmetry > 0,28855"], "Lá: 18 lượt bootstrap, tất cả nhãn ác.", CORAL),
    (6.88, "CÂY 2 → LÀNH", ["mean concave points ≤ 0,04989", "worst symmetry > 0,16385", "worst perimeter ≤ 127,20000"], "Lá: 256/259 lượt nhãn lành (98,84%).", TEAL),
]:
    panel(s, x, 2.12, 5.78, 3.93, BG)
    label(s, x + .24, 2.46, title, col, 5.3)
    for j, value in enumerate(lines):
        txt(s, x + .25, 3.17 + j * .62, 5.27, .42, f"{j+1}. {value}", 17.2, NAVY)
    txt(s, x + .25, 5.33, 5.27, .55, final, 15.4, col, True)
takeaway(s, "Các lượt bootstrap có thể lặp hàng; độ thuần tại lá không bảo đảm đúng với quan sát mới.")
note(s, "Các ngưỡng trên slide đã làm tròn để dễ đọc; JSON lưu giá trị đầy đủ. Mẫu có worst radius = 16,25; worst concave points = 0,14890; worst symmetry = 0,29620; mean concave points = 0,04819 và worst perimeter = 109,10. Theo hai chuỗi điều kiện khác nhau, cây thứ nhất đến lá chứa 18 lượt nhãn ác; cây thứ hai đến lá có 256 lượt lành và 3 lượt ác. Bootstrap có thể lặp hàng nên số lượt chưa phải số bệnh nhân khác nhau. Hai đường đi chỉ giải thích hai phiếu; tổ hợp cuối dùng đủ 51 cây.")

# 22 — Backup: calculations.
s = slide("B. Tính chỉ số với ác tính là lớp dương", "Phụ lục / Kiểm tra phép tính", "Bagging seed 42 · TP=49, FN=4, FP=3, TN=87", backup=True)
table(s, .67, 2.02, [3.1, 4.35, 4.5], [
    ["Chỉ số", "Công thức", "Kết quả"],
    ["Accuracy", "(TP + TN) / N", "136/143 = 95,10%"],
    ["Precision ác", "TP / (TP + FP)", "49/52 = 94,23%"],
    ["Recall ác", "TP / (TP + FN)", "49/53 = 92,45%"],
    ["F1 ác", "2TP / (2TP + FP + FN)", "98/105 = 93,33%"],
    ["Specificity", "TN / (TN + FP)", "87/90 = 96,67%"],
    ["Balanced accuracy", "Trung bình recall hai lớp", "94,56%"],
], .56, 17.5)
takeaway(s, "Accuracy test, balanced accuracy CV và OOB là các phép đánh giá khác nhau.")
note(s, "TP và FN đọc từ hàng nhãn thật ác tính; FP và TN từ hàng nhãn thật lành tính. Với Bagging tại seed 42, TP = 49, FN = 4, FP = 3 và TN = 87. Balanced accuracy bằng (49/53 + 87/90)/2 ≈ 94,56%. Vì lớp dương có nhãn 0, các hàm tính precision, recall và F1 dùng pos_label = 0. Phân biệt accuracy test 95,10% với balanced accuracy CV 95,26% trên train: chúng thuộc hai quy trình khác nhau. Tỷ lệ phiếu cây chưa phải xác suất bệnh đã hiệu chỉnh.")

# 23 — Backup: theory beyond the measured models.
s = slide("C. Boosting và Stacking khác Bagging ở đâu?", "Phụ lục / Mở rộng lý thuyết", "[3], [6], [7] · Chưa chạy thực nghiệm hai phương pháp này", backup=True)
for x, title, body_ in [(.68, "BOOSTING", "Học một chuỗi mô hình.\nAdaBoost tăng trọng số tương đối\ncho mẫu dự đoán sai.\nGradient Boosting giảm hàm mất mát."),
                        (6.88, "STACKING", "Học mô hình tầng hai.\nĐầu vào là dự đoán ngoài fold (OOF)\ncủa các mô hình cơ sở.\nTránh dùng dự đoán in-sample để học tầng hai.")]:
    panel(s, x, 2.2, 5.78, 3.75, BG)
    txt(s, x + .26, 2.59, 5.2, .5, title, 26, NAVY, True)
    txt(s, x + .26, 3.57, 5.2, 1.9, body_, 19, GRAY, leading=1.4)
takeaway(s, "Bagging đổi mẫu train để học các mô hình độc lập; hai hướng trên thay đổi cách phối hợp.")
note(s, "Trong ví dụ AdaBoost nhị phân, sai số có trọng số bằng 0,2 cho alpha = 0,5 ln 4 ≈ 0,693. Nhóm sai nhân 2, nhóm đúng nhân 0,5; sau chuẩn hóa, nhóm sai mang 50% trọng số. Đây chưa phải accuracy tăng từ 20% lên 50%. Stacking dùng dự đoán ngoài fold để mỗi quan sát train nhận dự đoán từ mô hình chưa học nó, rồi học meta-model. Khi dự đoán dữ liệu mới, các mô hình cơ sở học lại trên toàn train. Nếu hỏi 51 cây: đó là cấu hình ban đầu, khảo sát bổ sung chưa chứng minh tối ưu tổng quát.")

# 24 — Backup: sources and executable commands.
s = slide("D. Tài liệu tham khảo và cách tái lập", "Phụ lục / Nguồn", "Danh mục đầy đủ và DOI nằm trong báo cáo Word/PDF", backup=True)
references = [
    "[1] Breiman (1996). Bagging Predictors. Machine Learning 24, 123–140.",
    "[2] Breiman (2001). Random Forests. Machine Learning 45, 5–32.",
    "[3] scikit-learn User Guide. Ensemble methods.",
    "[4] scikit-learn API. load_breast_cancer.",
    "[5] scikit-learn User Guide. Decision Trees.",
    "[6] Freund & Schapire (1997). An Application to Boosting. JCSS 55(1).",
    "[7] Wolpert (1992). Stacked Generalization. Neural Networks 5(2).",
]
txt(s, .73, 2.03, 11.9, 2.6, "\n".join(references), 15.8, INK, leading=1.52)
panel(s, .67, 5.01, 11.95, 1.45, NAVY)
txt(s, .93, 5.22, 11.4, 1.05, "python de_5_ensemble_learning.py\npython experiments/report_extensions.py\npython experiments/controlled_comparison.py", 17.3, WHITE, font=MONO, leading=1.28)
note(s, "Tài liệu: [1] https://doi.org/10.1007/BF00058655; [2] https://doi.org/10.1023/A:1010933404324; [3] https://scikit-learn.org/stable/modules/ensemble.html; [4] https://scikit-learn.org/stable/modules/generated/sklearn.datasets.load_breast_cancer.html; [5] https://scikit-learn.org/stable/modules/tree.html; [6] https://doi.org/10.1006/jcss.1997.1504; [7] https://doi.org/10.1016/S0893-6080(05)80023-1. Mã nhóm: https://github.com/Huyngaongow/hocmaygk. Dùng Python 3.12 và requirements.txt. Chạy thí nghiệm gốc trước hai phân tích bổ sung. Kiểm thử: python -m unittest discover -s tests -v. Ví dụ: python examples/toy_bagging.py. Demo nhanh: python de_5_ensemble_learning.py --repeats 1 --output-dir outputs/quick. Dùng Python trong venv đã cài thư viện; trên Windows có thể gọi trực tiếp .venv\\Scripts\\python.exe.")

assert len(P.slides) == len(NOTES) == TOTAL_SLIDES
OUT = ROOT / "De_5_Ensemble_Learning.pptx"
P.save(OUT)
md = ["# Lời thuyết trình — Đề 5: Ensemble Learning", "",
      "24 slide: 20 slide chính và 4 slide dự phòng. Gợi ý luyện tập 15–18 phút; thời lượng thực tế theo giảng viên.", "",
      "Phân chia gợi ý: Nguyễn Nghĩa Hiếu slide 1–11; Trần Quốc Bảo slide 12–20; cả hai tham gia hỏi đáp. Đây là phương án luyện tập, không phải xác nhận lịch sử đóng góp.", "",
      "Nguồn số liệu: outputs/results.json, outputs/extensions/analysis.json và controlled_comparison.json. Các khảo sát bổ sung chỉ dùng train seed 42.", ""]
for idx, title, content in NOTES:
    md.extend([f"## Slide {idx:02d} — {title}", "", content, ""])
(ROOT / "presentation/loi_thuyet_trinh.md").write_text("\n".join(md), encoding="utf-8")
print(f"Saved {OUT} ({len(P.slides)} slides; editable shapes, chart, speaker notes and manual Fade transitions)")
