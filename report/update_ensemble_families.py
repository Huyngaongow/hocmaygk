"""Integrate the reproducible four-family extension into the faculty-form report.

Run after experiments/ensemble_families.py, using Python with python-docx.
The manual report is preserved; the named extension sections are updated in place.
PDF conversion is deliberately separate so the caller can inspect the result.
"""

from copy import deepcopy
import argparse
import json
from pathlib import Path
import re
import unicodedata
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report/Bao_cao_Ensemble_Learning_da_hoan_thien.docx"
DATA = ROOT / "outputs/extensions/family_comparison.json"
SUPPLEMENT = ROOT / "report/Mo_rong_Boosting_Voting_Stacking.md"


def text(paragraph, value):
    """Replace visible text while retaining paragraph and bookmark formatting."""
    keep = {qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")}
    for child in list(paragraph._p):
        if child.tag not in keep:
            paragraph._p.remove(child)
    run = paragraph.add_run(value)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    if paragraph.style.name.startswith("Heading"):
        run.bold = True
        run.font.size = Pt(16 if paragraph.style.name == "Heading 1" else 14)
    if paragraph.style.name == "Caption":
        run.font.size = Pt(12)


def find(doc, prefix):
    # Cached contents entries also start with the heading/caption text.
    return next(p for p in doc.paragraphs if p.text.startswith(prefix) and "\t" not in p.text)


def remove_section(doc, prefix, stop):
    starts = [p for p in doc.paragraphs if p.text.startswith(prefix)
              and p.style.name.startswith("Heading")]
    if not starts:
        return
    node = starts[0]._p
    while node is not None:
        nxt = node.getnext()
        if node.tag == qn("w:p") and Paragraph(node, doc).text.startswith(stop):
            break
        node.getparent().remove(node)
        node = nxt


def insert_paragraph(anchor, value, style="Nội dung văn bản"):
    node = OxmlElement("w:p")
    anchor._p.addprevious(node)
    p = Paragraph(node, anchor._parent)
    p.style = style
    text(p, value)
    pf = p.paragraph_format
    pf.space_after = Pt(5)
    pf.line_spacing = 1.5
    pf.keep_with_next = style.startswith("Heading") or style == "Caption"
    pf.keep_together = True if style.startswith("Heading") else False
    if style.startswith("Heading"):
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(6)
        # The faculty template numbers Heading styles automatically; these
        # headings already include their explicit chapter/section numbers.
        numpr = OxmlElement("w:numPr")
        numid = OxmlElement("w:numId"); numid.set(qn("w:val"), "0")
        numpr.append(numid); p._p.get_or_add_pPr().append(numpr)
    elif style == "Caption":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
        pf.line_spacing = 1
    else:
        pf.first_line_indent = Cm(1.27)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if value.startswith(("εₜ =", "wᵢ mới", "H(x) =")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Cm(0)
            pf.line_spacing = 1.15
            for run in p.runs:
                run.font.name = "Cambria Math"
            if value.startswith("H(x) = argmax_c"):
                text(p, "")
                for part, subscript in (("H(x) = argmax", False), ("c", True),
                                        (value.split("argmax_c", 1)[1], False)):
                    run = p.add_run(part)
                    run.font.name = "Cambria Math"; run.font.size = Pt(13)
                    run.font.subscript = subscript
    return p


def insert_table(doc, anchor, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for cell, label in zip(t.rows[0].cells, headers):
        cell.text = label
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            if i == 0:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E7EDF2")
                cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_with_next = False
                for r in p.runs:
                    r.font.name = "Times New Roman"; r.font.size = Pt(11.5)
                    r.bold = i == 0
    anchor._p.addprevious(t._tbl)
    return t


def insert_code(doc, anchor, source):
    t = insert_table(doc, anchor, [source], [], [15.5])
    for p in t.cell(0, 0).paragraphs:
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = Pt(11.5)
        for r in p.runs:
            r.bold = False; r.font.name = "DejaVu Sans Mono"; r.font.size = Pt(10)
    return t


def add_heading(anchor, value, level=2):
    return insert_paragraph(anchor, value, "Heading " + str(level))


def fmt(v):
    return f"{100 * v:.2f}".replace(".", ",")


def avg(row, metric):
    return f"{fmt(row[metric + '_mean'])} ± {fmt(row[metric + '_std'])}"


NAMES = {
    "Decision Tree": "Cây đơn (độ sâu CV)",
    "Manual Bagging": "Manual Bagging",
    "Random Forest (sklearn)": "Random Forest",
    "Decision Tree (depth=3)": "Cây đơn (độ sâu 3)",
    "Logistic Regression": "Logistic Regression",
    "Gaussian NB": "Gaussian NB",
    "AdaBoost": "AdaBoost",
    "Voting (hard)": "Voting hard",
    "Voting (soft)": "Voting soft",
    "Stacking": "Stacking",
}


def rebuild_contents(doc):
    """Linked lists of headings/captions with Word/LibreOffice PAGEREF fields."""
    toc_title = find(doc, "MỤC LỤC")
    abbr = next(t for t in doc.tables if t.cell(0, 0).text == "Chữ viết tắt")
    node = toc_title._p.getnext()
    while node is not None and node is not abbr._tbl:
        nxt = node.getnext(); node.getparent().remove(node); node = nxt
    # Remove stale contents bookmarks; the new stable list replaces their links.
    for el in list(doc._element.xpath(".//w:bookmarkStart")):
        if el.get(qn("w:name"), "").startswith(("Report_", "EnsReport_")):
            bid = el.get(qn("w:id"))
            for end in list(doc._element.xpath(".//w:bookmarkEnd")):
                if end.get(qn("w:id")) == bid:
                    end.getparent().remove(end)
            el.getparent().remove(el)
    marker = Paragraph(OxmlElement("w:p"), doc._body)
    abbr._tbl.addprevious(marker._p)
    main_start = find(doc, "CHƯƠNG 1.")
    main = []
    after = False
    for p in doc.paragraphs:
        if p._p is main_start._p: after = True
        if after and p.style.name in ("Heading 1", "Heading 2", "Heading 3", "Caption"):
            main.append(p)
    fronts = []
    for title in ("DANH MỤC HÌNH VẼ", "DANH MỤC BẢNG BIỂU", "DANH MỤC THUẬT TOÁN", "DANH MỤC CHỮ VIẾT TẮT"):
        p = insert_paragraph(marker, title, "Heading 1")
        p.paragraph_format.page_break_before = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fronts.append(p)
    entries = fronts + main
    marks = {}
    for i, p in enumerate(entries, 6001):
        name = f"EnsReport_{i}"
        start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(i)); start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(i))
        p._p.insert(1 if p._p.pPr is not None else 0, start); p._p.append(end)
        marks[id(p._p)] = name
    def item(anchor, target, level=0):
        p = insert_paragraph(anchor, "", "Normal")
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(level * .4)
        p.paragraph_format.keep_together = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tabs = p._p.get_or_add_pPr().find(qn("w:tabs"))
        if tabs is not None: p._p.get_or_add_pPr().remove(tabs)
        tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab")
        tab.set(qn("w:pos"), "8674"); tab.set(qn("w:val"), "right"); tab.set(qn("w:leader"), "dot")
        tabs.append(tab); p._p.get_or_add_pPr().append(tabs)
        link = OxmlElement("w:hyperlink"); link.set(qn("w:anchor"), marks[id(target._p)])
        run = OxmlElement("w:r"); pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), "Times New Roman"); fonts.set(qn("w:hAnsi"), "Times New Roman")
        pr.append(fonts); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "24"); pr.append(sz)
        run.append(pr); tx = OxmlElement("w:t"); tx.text = target.text; run.append(tx); link.append(run); p._p.append(link)
        p.add_run("\t")
        number_format = "roman" if target in fronts else "Arabic"
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), f" PAGEREF {marks[id(target._p)]} \\h \\* {number_format} ")
        rr = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "1"; rr.append(tt); field.append(rr); p._p.append(field)
        return p
    for p in fronts + [p for p in main if p.style.name.startswith("Heading")]:
        level = int(p.style.name[-1]) - 1
        item(fronts[0], p, level)
    for p in main:
        if p.style.name == "Caption":
            if p.text.startswith("Hình "): item(fronts[1], p)
            elif p.text.startswith("Bảng "): item(fronts[2], p)
            elif p.text.startswith("Thuật toán "): item(fronts[3], p)
    marker._p.getparent().remove(marker._p)
    settings = doc.settings.element
    for old in list(settings.findall(qn("w:updateFields"))): settings.remove(old)
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)


def refresh_pages(pdf_path):
    """Store real page numbers so Word displays a complete contents on opening."""
    doc = Document(REPORT)
    def norm(s):
        return "".join(unicodedata.normalize("NFKC", s).split())
    pages = subprocess.check_output(["pdftotext", "-raw", str(pdf_path), "-"]).decode("utf-8").split("\f")
    main_page = next(i for i, s in enumerate(pages)
                     if re.search(r"^1\s*\nCHƯƠNG 1\.", s))
    targets = {}
    for p in doc.paragraphs:
        marks = p._p.findall(qn("w:bookmarkStart"))
        marks = [m for m in marks if m.get(qn("w:name"), "").startswith("EnsReport_")]
        if not marks:
            continue
        is_front = p.text.startswith("DANH MỤC")
        search = range(main_page) if is_front else range(main_page, len(pages))
        candidates = []
        for i in search:
            s = pages[i]
            if is_front:
                lines = [x.strip() for x in s.splitlines() if x.strip()]
                matched = len(lines) > 1 and norm(lines[1]) == norm(p.text)
            else:
                matched = norm(p.text) in norm(s)
            if matched:
                first = next(x.strip() for x in s.splitlines() if x.strip())
                if re.fullmatch(r"[ivxlcdm]+|\d+", first):
                    candidates.append((i, first))
        if not candidates:
            raise ValueError(f"Không tìm được trang của {p.text!r}")
        for mark in marks:
            targets[mark.get(qn("w:name"))] = candidates[0][1]
    count = 0
    for field in doc._element.xpath(".//w:fldSimple"):
        instr = field.get(qn("w:instr"), "")
        m = re.search(r"PAGEREF\s+(EnsReport_\d+)", instr)
        if not m:
            continue
        for tx in field.iter(qn("w:t")):
            tx.text = targets[m.group(1)]
        field.set(qn("w:dirty"), "false")
        count += 1
    doc.save(REPORT)
    print(f"Updated {count} cached page references from {pdf_path}")


def export_pdf():
    """Render twice, preserving native Word fields and reliable Roman PDF lists."""
    output = REPORT.with_suffix(".pdf")
    with tempfile.TemporaryDirectory(prefix="hocmaygk-family-report-") as temporary:
        folder = Path(temporary)
        first = folder / "first"; first.mkdir()
        def convert(source, destination, profile):
            subprocess.run([shutil.which("soffice"), f"-env:UserInstallation={profile.as_uri()}",
                            "--headless", "--convert-to", "pdf", "--outdir", str(destination),
                            str(source)], check=True, timeout=120)
            path = destination / source.with_suffix(".pdf").name
            if not path.is_file(): raise RuntimeError("PDF conversion did not produce an output")
            return path
        first_pdf = convert(REPORT, first, folder / "profile-first")
        refresh_pages(first_pdf)
        # LibreOffice ignores Roman formatting on PAGEREF in this Word template.
        # Flatten these four references only in an export copy. The editable
        # final Word keeps all linked fields, including their correct cache.
        copy = Document(REPORT)
        for field in list(copy._element.xpath(".//w:fldSimple")):
            if "\\* roman" in field.get(qn("w:instr"), ""):
                for run in list(field): field.addprevious(deepcopy(run))
                field.getparent().remove(field)
        source = folder / REPORT.name; copy.save(source)
        final_dir = folder / "final"; final_dir.mkdir()
        final_pdf = convert(source, final_dir, folder / "profile-final")
        shutil.copyfile(final_pdf, output)
        refresh_pages(output)
    print(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--refresh-pages", type=Path, help="Refresh cached page references from the rendered PDF only")
    parser.add_argument("--export-pdf", action="store_true", help="Build, render and verify cached contents pages")
    args = parser.parse_args()
    if args.refresh_pages:
        refresh_pages(args.refresh_pages)
        return
    data = json.loads(DATA.read_text(encoding="utf-8"))
    rows = {r["model"]: r for r in data["summary"]}
    if set(rows) != set(NAMES) or any(r["n_runs"] != 10 for r in rows.values()):
        raise SystemExit("Cần family_comparison.json đầy đủ 10 mô hình × 10 seed.")
    doc = Document(REPORT)
    for start, stop in (("2.9 ", "CHƯƠNG 3."), ("4.10 ", "CHƯƠNG 5."), ("5.16 ", "CHƯƠNG 6.")):
        remove_section(doc, start, stop)
    # Scope is consistent throughout the existing manual report.
    replacements = {
        "Mục tiêu tổng quát là tìm hiểu và triển khai cơ chế Bagging, đồng thời đánh giá phương pháp trong một quy trình thực nghiệm có thể kiểm tra lại.":
        "Mục tiêu tổng quát là tìm hiểu Ensemble Learning, triển khai sâu cơ chế Bagging và minh họa AdaBoost, Voting, Stacking trong một quy trình thực nghiệm có thể kiểm tra lại.",
        "Đối tượng nghiên cứu là Bagging sử dụng cây quyết định cho bài toán phân loại.":
        "Đối tượng nghiên cứu là Ensemble Learning cho phân loại, trong đó Bagging cây quyết định là trọng tâm triển khai và tính toán minh họa; AdaBoost, Voting và Stacking được bổ sung để đối chiếu các cơ chế tổ hợp.",
        "Boosting và Stacking được giới thiệu để cung cấp bức tranh tổng quan, chưa được triển khai trong phần đối chiếu thực nghiệm. Bài làm mang mục đích minh họa phương pháp học máy, chưa đánh giá khả năng sử dụng trong chẩn đoán lâm sàng.":
        "AdaBoost, Voting hard/soft và Stacking được triển khai minh họa bằng scikit-learn trên cùng các lần chia dữ liệu; các cấu hình được công bố ở mục 4.10. Bài làm mang mục đích học thuật, chưa đánh giá khả năng chẩn đoán lâm sàng.",
    }
    for p in doc.paragraphs:
        if p.text in replacements: text(p, replacements[p.text])
        elif p.text == "BÁO CÁO GIỮA KỲ":
            text(p, "TIỂU LUẬN GIỮA KỲ"); p.runs[0].bold = True; p.runs[0].font.size = Pt(16)
        elif p.text == "ĐỀ 5 — ENSEMBLE LEARNING\nBAGGING CÂY QUYẾT ĐỊNH":
            text(p, "ENSEMBLE LEARNING"); p.runs[0].bold = True; p.runs[0].font.size = Pt(20)
        elif p.text == "Sử dụng năm quan sát giả lập một chiều": text(p, p.text + ".")
        elif "Tính tuyến tính của kỳ vọng cho [1]:" in p.text:
            text(p, p.text.replace("Tính tuyến tính của kỳ vọng cho [1]:", "Theo tính tuyến tính của kỳ vọng:"))
        elif p.text.startswith("Bộ kiểm thử hiện có 26 kiểm thử"):
            text(p, "Bộ kiểm thử tự động kiểm tra các bước cốt lõi của Bagging, tổng hợp hard/soft, ánh xạ xác suất khi cây chỉ học một lớp, phép đo bất đồng và giao thức so sánh các họ ensemble. Chạy toàn bộ bộ kiểm thử trước khi tái lập kết quả.")
        elif "Mã nguồn cùng cấu hình, kết quả và 26 kiểm thử" in p.text:
            text(p, p.text.replace("26 kiểm thử", "bộ kiểm thử tự động"))
    # Add concise mechanism descriptions without expanding the Bagging hand example.
    anchor = find(doc, "CHƯƠNG 3.")
    add_heading(anchor, "2.9 Boosting: minh họa bằng AdaBoost")
    ada_text = [
        "AdaBoost xây dựng tuần tự các mô hình cơ sở. Mỗi vòng học sử dụng trọng số của các quan sát; mẫu bị dự đoán sai được tăng trọng số tương đối để vòng sau chú ý hơn. Phần minh họa sử dụng cây độ sâu 1 (decision stump), tối đa 51 vòng và learning_rate = 1 [3], [6].",
        "Quy trình SAMME trong bài toán hai lớp: (1) khởi tạo wᵢ = 1/N; (2) học cây hₜ với trọng số wᵢ; (3) tính tổng trọng số các mẫu sai εₜ; (4) tính trọng số phiếu αₜ; (5) tăng trọng số mẫu sai và chuẩn hóa; (6) lặp lại rồi tổng hợp phiếu có trọng số. Với 0 < εₜ < 0,5, các công thức cho bước cập nhật là:",
        "εₜ = ∑ᵢ wᵢ I[hₜ(xᵢ) ≠ yᵢ];    αₜ = ln((1 − εₜ)/εₜ).",
        "wᵢ mới ∝ wᵢ × exp(αₜ I[hₜ(xᵢ) ≠ yᵢ]).",
        "H(x) = argmax_c ∑ₜ αₜ I[hₜ(x) = c].",
        "Đây là quy ước SAMME của AdaBoostClassifier trong scikit-learn 1.8 cho hai lớp, không dùng trực tiếp hệ số ½ của một cách viết AdaBoost nhị phân khác. Cây hoàn hảo có thể làm thuật toán dừng sớm; vì vậy 51 là số vòng tối đa, không luôn là số cây thực học.",
        "Ví dụ: năm mẫu có trọng số bằng nhau 0,2; cây đầu dự đoán sai một mẫu nên ε = 0,2 và α = ln 4 ≈ 1,386. Mẫu sai nhân 4 thành 0,8; bốn mẫu đúng giữ 0,2 mỗi mẫu. Tổng trọng số chưa chuẩn hóa là 1,6. Sau chuẩn hóa, mẫu sai nhận 0,5; mỗi mẫu đúng nhận 0,125. Vòng kế tiếp tập trung hơn vào mẫu sai, không có nghĩa accuracy đã tăng lên 50%.",
    ]
    for s in ada_text: insert_paragraph(anchor, s)
    add_heading(anchor, "2.10 Voting với các mô hình khác loại")
    voting_text = [
        "Voting trong phần mở rộng kết hợp ba mô hình khác loại: cây quyết định độ sâu 3, Logistic Regression và Gaussian Naive Bayes. Cây chia không gian bằng điều kiện; Logistic Regression học ranh giới tuyến tính; Gaussian NB mô hình hóa phân bố đặc trưng theo lớp. Sự khác nhau về giả thiết học có thể tạo ra các dự đoán bổ sung cho nhau, nhưng không bảo đảm tổ hợp luôn tốt hơn từng mô hình.",
        "Quy trình: học riêng ba mô hình trên train; lấy nhãn cho hard voting hoặc xác suất lớp cho soft voting; tổng hợp với trọng số bằng nhau; đánh giá trên test. Công thức ở mục 2.5 vẫn áp dụng. Hard và soft sử dụng cùng cấu hình mô hình cơ sở để việc đối chiếu không thay đổi đồng thời cả kiến trúc. Soft voting có thể bị ảnh hưởng bởi chất lượng xác suất của từng mô hình; bài chưa hiệu chuẩn xác suất [3].",
    ]
    for s in voting_text: insert_paragraph(anchor, s)
    add_heading(anchor, "2.11 Stacking và dự đoán ngoài fold")
    stack_text = [
        "Stacking học một mô hình tầng hai để kết hợp đầu ra của các mô hình cơ sở [7]. Ba mô hình cơ sở giống phần Voting; mô hình tầng hai là Logistic Regression. Khác với Voting dùng quy tắc cố định, Stacking học cách kết hợp từ dữ liệu.",
        "Quy trình trên train: (1) chia thành năm fold phân tầng; (2) trong từng fold, học các mô hình cơ sở trên bốn fold còn lại và dự đoán xác suất cho fold được giữ; (3) ghép các dự đoán ngoài fold (OOF) thành đầu vào tầng hai; (4) học tầng hai trên các đầu ra OOF và nhãn train; (5) học lại các mô hình cơ sở trên toàn train. Khi gặp test, ba mô hình cơ sở tạo đầu vào cho tầng hai dự đoán nhãn cuối.",
        "Trong scikit-learn, với hai lớp và stack_method='predict_proba', một cột xác suất dư thừa được bỏ ở mỗi mô hình cơ sở. Ba cột còn lại tương ứng xác suất lớp 1 (lành tính), không phải mặc định là xác suất ác tính. StandardScaler nằm trong Pipeline của Logistic Regression nên mỗi lần fit chỉ học thống kê từ phần dữ liệu được phép học. Không dùng dự đoán in-sample để thay thế OOF; tập test không tham gia học tầng hai [3], [7].",
    ]
    for s in stack_text: insert_paragraph(anchor, s)
    # Application protocol: deliberately descriptive, with declared fixed configs.
    anchor = find(doc, "CHƯƠNG 5.")
    add_heading(anchor, "4.10 Thiết kế so sánh các hướng Ensemble Learning")
    protocol_text = [
        "Phần mở rộng dùng đúng mười lần chia phân tầng 75/25, seed 42–51 như thí nghiệm gốc. Trong mỗi seed, mọi mô hình dùng chung 426 quan sát train và 143 quan sát test. Ba kết quả gốc được giữ để đối chiếu; bảy cấu hình mới được khai báo trước lần chạy bổ sung và không được điều chỉnh theo điểm test.",
        "Bộ cơ sở của Voting/Stacking gồm cây độ sâu 3 cố định, Logistic Regression có StandardScaler trong Pipeline và Gaussian NB. Ba mô hình này được đánh giá riêng để kiểm tra lợi ích của tổ hợp. Cây độ sâu 3 được ghi thành một dòng riêng, vì khác quy trình chọn độ sâu bằng CV của cây đơn trong thí nghiệm gốc. Cố định kiến trúc trước khi tạo OOF giúp không dùng fold được giữ để lựa chọn cấu hình cây.",
    ]
    for s in protocol_text: insert_paragraph(anchor, s)
    insert_paragraph(anchor, "Bảng 4.6. Cấu hình phần mở rộng các họ ensemble", "Caption")
    config_rows = [
        ["AdaBoost", "Tối đa 51 cây độ sâu 1; learning_rate = 1; SAMME"],
        ["Voting hard/soft", "Cây độ sâu 3 + Logistic Regression chuẩn hóa + Gaussian NB; trọng số bằng nhau"],
        ["Stacking", "Cùng ba mô hình; OOF 5 fold phân tầng có trộn; predict_proba; tầng hai Logistic Regression; passthrough=False"],
        ["Đối chứng cơ sở", "Đánh giá riêng cây độ sâu 3, Logistic Regression và Gaussian NB"],
    ]
    insert_table(doc, anchor, ["Nhóm", "Cấu hình cố định"], config_rows, [3.8, 11.7])
    protocol_end = "Đây là khảo sát bổ sung sau khi đã xem kết quả thí nghiệm gốc, nên được diễn giải như đối chiếu mô tả, không phải xác nhận độc lập trên dữ liệu mới. Các mô hình chưa được tối ưu riêng với cùng ngân sách; bảng kết quả so sánh các cấu hình đã công bố, không tách riêng tác động nhân quả của tên phương pháp. Trung bình và SD mẫu dùng cùng quy ước với mục 4.7."
    insert_paragraph(anchor, protocol_end)
    # Results from the actual generated JSON, with no numbers typed by hand.
    anchor = find(doc, "CHƯƠNG 6.")
    add_heading(anchor, "5.16 Triển khai và kết quả của AdaBoost, Voting, Stacking")
    add_heading(anchor, "5.16.1 Tổ chức triển khai", 3)
    code_intro = "Tệp experiments/ensemble_families.py dựng các cấu hình bằng scikit-learn, lưu chỉ số của từng seed và tổng hợp kết quả. Đoạn mã sau minh họa cách gọi các mô hình từ hàm build_models; đầy đủ cấu hình và cách tạo OOF nằm trong mã nguồn đi kèm."
    intro_paragraph = insert_paragraph(anchor, code_intro)
    intro_paragraph.paragraph_format.keep_with_next = True
    source = """from sklearn.datasets import load_breast_cancer
from sklearn.model_selection import train_test_split
from experiments.ensemble_families import build_models

data = load_breast_cancer()
X_train, X_test, y_train, y_test = train_test_split(
    data.data, data.target, test_size=0.25,
    stratify=data.target, random_state=42,
)
models = build_models(seed=42)
for name in ("AdaBoost", "Voting (hard)",
             "Voting (soft)", "Stacking"):
    model = models[name]
    model.fit(X_train, y_train)
    prediction = model.predict(X_test)
    print(name, (prediction == y_test).mean())"""
    insert_code(doc, anchor, source)
    insert_paragraph(anchor, "Chỉ tiêu precision, recall và F1 ác tính trong chương trình dùng pos_label=0; ma trận nhầm lẫn giữ thứ tự [0, 1]. Scikit-learn thực hiện vòng lặp AdaBoost, Voting và OOF Stacking; nhóm tự tổ chức chia dữ liệu, cấu hình, đánh giá và xuất bằng chứng. Phần Bagging vẫn là phần cơ chế tổ hợp do nhóm tự triển khai.")
    add_heading(anchor, "5.16.2 Kết quả trên mười lần chia", 3)
    insert_paragraph(anchor, "Bảng 5.16. Kết quả so sánh các cấu hình trên mười holdout", "Caption")
    metric_rows = [[NAMES[name], avg(rows[name], "accuracy"), avg(rows[name], "balanced_accuracy"), avg(rows[name], "recall_malignant")]
                   for name in NAMES]
    metrics_table = insert_table(doc, anchor, ["Mô hình", "Accuracy", "Balanced accuracy", "Recall ác"], metric_rows, [5.0, 3.5, 3.5, 3.5])
    for i, row in enumerate(metrics_table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = i < len(metrics_table.rows) - 1
    units_note = insert_paragraph(anchor, "Trung bình tính theo %, SD mẫu tính theo điểm phần trăm (ddof=1). Mỗi dòng có mười seed cùng các tập train/test tương ứng. Nguồn: outputs/extensions/family_comparison.json và family_comparison_summary.csv; F1, precision, specificity và số đếm từng seed được lưu trong tệp kết quả.")
    units_note.paragraph_format.keep_together = True
    units_note.paragraph_format.line_spacing = 1.05
    units_note.paragraph_format.first_line_indent = Cm(0)
    units_note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in units_note.runs: run.font.size = Pt(11)
    baseline = max((rows[n] for n in ("Decision Tree (depth=3)", "Logistic Regression", "Gaussian NB")), key=lambda r: r["accuracy_mean"])
    top = max(rows.values(), key=lambda r: r["accuracy_mean"])
    discussion = [
        f"Trong toàn bộ cấu hình khảo sát, {NAMES[top['model']]} đạt accuracy trung bình cao nhất ({fmt(top['accuracy_mean'])}%). Kết quả này cho thấy các ensemble được xét chưa vượt qua mô hình đơn mạnh nhất. Đây là thứ hạng mô tả của lần khảo sát, chưa phải bằng chứng phương pháp đó luôn tốt nhất.",
        f"Voting hard đạt accuracy trung bình {fmt(rows['Voting (hard)']['accuracy_mean'])}%, Voting soft đạt {fmt(rows['Voting (soft)']['accuracy_mean'])}%, Stacking đạt {fmt(rows['Stacking']['accuracy_mean'])}% và AdaBoost đạt {fmt(rows['AdaBoost']['accuracy_mean'])}%. Cần đọc cùng recall ác tính và mức dao động; kết hợp thêm mô hình hoặc học thêm tầng hai không tự động bảo đảm cải thiện so với mô hình cơ sở mạnh.",
        "Bagging học các mô hình trên bootstrap; AdaBoost cập nhật trọng số qua các vòng; Voting áp dụng quy tắc kết hợp cố định; Stacking học tầng kết hợp từ OOF. Các cấu hình sử dụng loại mô hình và chi phí khác nhau, vì vậy bảng minh họa hành vi của các hướng tiếp cận thay vì chứng minh một cơ chế tổ hợp vượt trội trong mọi điều kiện.",
        "Các holdout có thể giao nhau; SD chỉ mô tả dao động. Khảo sát này chưa hiệu chuẩn xác suất, kiểm định thống kê hay đánh giá trên dữ liệu ngoài độc lập.",
    ]
    for s in discussion: insert_paragraph(anchor, s)
    # Existing summary, scope table, code/file index and reproduction instructions.
    for p in doc.paragraphs:
        if p.text.startswith("Báo cáo bổ sung khảo sát số cây, phân tích mẫu dự đoán sai"):
            text(p, "Báo cáo bổ sung khảo sát số cây, phân tích mẫu sai, đối chứng đặc trưng–quy tắc tổng hợp và minh họa AdaBoost, Voting, Stacking trên cùng mười lần chia dữ liệu. Các phép đối chiếu giúp phân biệt cơ chế tổ hợp; kết quả được diễn giải trong phạm vi cấu hình và dữ liệu đã khảo sát.")
        elif p.text.startswith("Từ khóa:"):
            text(p, "Từ khóa: Ensemble Learning, Bagging, AdaBoost, Voting, Stacking, bootstrap, OOB, OOF.")
        elif p.text.startswith("Các mục tiêu cụ thể gồm giải thích nguyên lý Ensemble Learning;") and "Phần mở rộng minh họa" not in p.text:
            text(p, p.text + " Phần mở rộng minh họa và đối chiếu AdaBoost, Voting, Stacking bằng thư viện scikit-learn.")
        elif p.text.startswith("Nhóm đã trình bày Ensemble Learning và tự triển khai bootstrap") and "AdaBoost, Voting và Stacking được triển khai bổ sung" not in p.text:
            text(p, p.text + " AdaBoost, Voting và Stacking được triển khai bổ sung bằng scikit-learn, có công bố cấu hình và kết quả trên cùng các lần chia dữ liệu.")
        elif p.text.startswith("Random Forest có accuracy trung bình cao hơn Bagging"):
            text(p, "Trong thí nghiệm gốc, " + p.text)
        elif p.text.startswith("Trong thí nghiệm gốc, random Forest"):
            text(p, p.text.replace("Trong thí nghiệm gốc, random Forest", "Trong thí nghiệm gốc, Random Forest", 1))
        elif p.text.startswith("Các mô hình đều dựa trên cây nên không yêu cầu chuẩn hóa"):
            text(p, p.text.replace("Các mô hình đều", "Ba mô hình trong thí nghiệm gốc đều", 1))
        elif p.text.startswith("Độ sâu được chọn bằng cây đơn rồi dùng chung, chưa tối ưu riêng từng mô hình."):
            value = "Trong thí nghiệm gốc, " + p.text[0].lower() + p.text[1:]
            if "Các cấu hình mới ở mục 4.10" not in value:
                value += " Các cấu hình mới ở mục 4.10 được cố định để minh họa, chưa tối ưu riêng hoặc so sánh theo cùng ngân sách tính toán; phần mở rộng được thực hiện sau khi đã quan sát kết quả gốc."
            text(p, value)
        elif p.text.startswith("Bagging có thể cải thiện so với một cây đơn bằng cách kết hợp nhiều mô hình"):
            text(p, "Bagging có thể cải thiện so với một cây đơn bằng cách kết hợp nhiều mô hình học trên bootstrap, nhưng không loại bỏ mọi sai sót. Các đối chiếu cho thấy cần xem đồng thời số cây, chất lượng mô hình cơ sở và quy tắc tổng hợp. Trong phần mở rộng, Logistic Regression đơn lẻ đạt các chỉ số trung bình cao nhất; kết quả này được giới hạn trong dữ liệu, cấu hình và giao thức đã khảo sát.")
        elif p.text.startswith("Giá trị của báo cáo là mối liên hệ kiểm chứng được"):
            text(p, "Giá trị của báo cáo là mối liên hệ kiểm chứng được giữa nguyên lý, thuật toán, ví dụ, mã nguồn và kết quả. Bagging được giải thích sâu; AdaBoost, Voting và Stacking mở rộng bức tranh học tổ hợp bằng những cơ chế khác nhau. Các kết luận được giới hạn trong giao thức đã công bố.")
        elif p.text.startswith("Hai tệp controlled_comparison.* trong Bảng A.1"):
            text(p, "Các tệp controlled_comparison.* và family_comparison.* trong Bảng A.1 nằm trong outputs/extensions/. Số liệu được lấy từ các tệp chương trình sinh; việc làm tròn có thể tạo sai khác ở chữ số cuối. Các nghiên cứu bổ sung không thay đổi kết quả gốc; khảo sát các họ ensemble là đối chiếu mô tả sau thí nghiệm gốc.")
    scope = next(t for t in doc.tables if t.cell(0, 0).text == "Phương pháp")
    for row in scope.rows[1:]:
        if row.cells[0].text == "Boosting": row.cells[2].text = "Minh họa AdaBoost"
        elif row.cells[0].text == "Voting": row.cells[2].text = "Quy tắc kết hợp và mô hình dị loại"
        elif row.cells[0].text == "Stacking": row.cells[2].text = "Minh họa OOF và tầng hai"
    index = next(t for t in doc.tables if t.cell(0, 0).text == "Thành phần" and "Chức năng" in t.cell(0, 1).text)
    if not any("ensemble_families" in row.cells[0].text for row in index.rows):
        cells = index.add_row().cells; cells[0].text = "experiments/ensemble_families.py"; cells[1].text = "So sánh AdaBoost, Voting và Stacking"
    appindex = next(t for t in doc.tables if t.cell(0, 0).text == "Tệp / thư mục")
    if not any("family_comparison" in row.cells[0].text for row in appindex.rows):
        for file, desc in (("family_comparison.json", "Cấu hình, split, OOF, chỉ số và dự đoán của mười seed"), ("family_comparison_summary.csv", "Bảng trung bình và SD để đối chiếu Bảng 5.16")):
            cells = appindex.add_row().cells; cells[0].text = file; cells[1].text = desc
    # Code spans use a fixed-width font; keep algorithm pseudo-code in its own style.
    for table in doc.tables:
        if len(table.columns) == 1 and len(table.rows) == 1:
            cell = table.cell(0, 0)
            if re.search(r"(^|\n)(import |from |self\.|def |python |h1:)", cell.text):
                if "python experiments/controlled_comparison.py" in cell.text and "ensemble_families.py" not in cell.text:
                    cell.text += "\npython experiments/ensemble_families.py"
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = Pt(11.5)
                    p.paragraph_format.first_line_indent = Cm(0)
                    for r in p.runs:
                        r.font.name = "DejaVu Sans Mono"; r.font.size = Pt(10)
        for row in table.rows:
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None: trpr.append(OxmlElement("w:cantSplit"))
    # A self-contained units note for the tree-count table.
    count_caption = find(doc, "Bảng 5.9.")
    node = count_caption._p.getnext()
    while node is not None and node.tag != qn("w:tbl"): node = node.getnext()
    after = node.getnext() if node is not None else None
    if after is not None:
        existing = Paragraph(after, doc).text if after.tag == qn("w:p") else ""
        if existing != "Trung bình tính theo %, SD tính theo điểm phần trăm.":
            ap = Paragraph(OxmlElement("w:p"), doc._body); node.addnext(ap._p)
            ap.style = "Nội dung văn bản"; text(ap, "Trung bình tính theo %, SD tính theo điểm phần trăm.")
            ap.paragraph_format.space_after = Pt(3); ap.paragraph_format.line_spacing = 1
    # Keep the prediction heading, caption and first algorithm block together.
    find(doc, "3.2 Thuật toán dự đoán").paragraph_format.keep_with_next = True
    rebuild_contents(doc)
    doc.core_properties.title = "Ensemble Learning — Bagging, AdaBoost, Voting và Stacking"
    doc.save(REPORT)
    # The same material is available without requiring Word editing by script.
    md = ["# Bổ sung AdaBoost, Voting và Stacking", "", "Các mục dưới đây đã được tích hợp vào báo cáo Word/PDF. Bagging vẫn là trọng tâm triển khai.", "", "## 2.9 AdaBoost", ""]
    md += [s + "\n" for s in ada_text]
    md += ["## 2.10 Voting", ""] + [s + "\n" for s in voting_text]
    md += ["## 2.11 Stacking", ""] + [s + "\n" for s in stack_text]
    md += ["## 4.10 Thiết kế thực nghiệm", ""] + [s + "\n" for s in protocol_text]
    md += ["| Nhóm | Cấu hình cố định |", "|---|---|"] + ["| " + " | ".join(r) + " |" for r in config_rows]
    md += ["", protocol_end, "", "## 5.16 Code và kết quả", "", code_intro, "", "```python", source, "```", "", "| Mô hình | Accuracy | Balanced accuracy | Recall ác |", "|---|---:|---:|---:|"]
    md += ["| " + " | ".join(r) + " |" for r in metric_rows]
    md += ["", "Trung bình tính theo %, SD mẫu theo điểm phần trăm, ddof=1; mười seed42–51.", ""]
    md += [s + "\n" for s in discussion]
    SUPPLEMENT.write_text("\n".join(md), encoding="utf-8")
    print(REPORT)
    print(SUPPLEMENT)
    if args.export_pdf:
        export_pdf()


if __name__ == "__main__":
    main()
