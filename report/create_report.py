"""Create the Vietnamese midterm paper from the actual experiment and toy example.

Run from any directory: .venv/bin/python report/create_report.py
Then export the DOCX with LibreOffice. --toc-pdf <previous-export.pdf> can
refresh the contents page with page numbers read from that actual PDF.
"""

import argparse
import json
from pathlib import Path
import re
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from examples.toy_bagging import build_toy_example

REPORT = ROOT / "report"
STEM = "Bao_cao_De_5_Ensemble_Learning"
TEAL = "087E83"
NAVY = "173246"
LABEL = {0: "ác tính", 1: "lành tính"}
MODEL = {"Decision Tree": "Cây quyết định đơn", "Manual Bagging": "Bagging tự triển khai",
         "Random Forest (sklearn)": "Random Forest (sklearn)"}
TOC_HEADINGS = ["MỞ ĐẦU", "1. NGUYÊN LÝ HOẠT ĐỘNG", "2. QUY TRÌNH VÀ VÍ DỤ TÍNH TAY",
                "3. BÀI TOÁN ỨNG DỤNG VÀ THIẾT KẾ THỰC NGHIỆM",
                "4. TRIỂN KHAI MÃ NGUỒN VÀ KẾT QUẢ", "KẾT LUẬN", "TÀI LIỆU THAM KHẢO"]


def pct(value, digits=2):
    return f"{100 * value:.{digits}f}%".replace(".", ",")


def number(value, digits=2):
    return f"{value:.{digits}f}".replace(".", ",")


def mean_sd(row, key):
    return f"{number(row[key + '_mean'] * 100)} ± {number(row[key + '_std'] * 100)}"


def field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tcpr.append(node)


class Paper:
    """One source of paragraphs/tables for both editable DOCX and Markdown."""

    def __init__(self):
        self.doc = Document()
        self.md = []
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.9)
        sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.0)
        sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)
        sec.different_first_page_header_footer = True
        styles = self.doc.styles
        for name in ["Normal", "Body Text", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
            styles[name].font.name = "Times New Roman"
            styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        normal = styles["Normal"]
        normal.font.size = Pt(12.5)
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(7)
        for name, size in [("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12.5)]:
            styles[name].font.size = Pt(size)
            styles[name].font.color.rgb = RGBColor.from_string(NAVY)
            styles[name].font.bold = True
            styles[name].paragraph_format.space_before = Pt(10)
            styles[name].paragraph_format.space_after = Pt(8)
        styles["Caption"].font.size = Pt(10.5)
        styles["Caption"].font.italic = True
        styles["Caption"].font.color.rgb = RGBColor.from_string("566A76")
        hp = sec.header.paragraphs[0]
        hp.text = "NHẬP MÔN HỌC MÁY   |   ĐỀ 5 — ENSEMBLE LEARNING"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in hp.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("566A76")
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field(fp, "PAGE")
        self.doc.core_properties.title = "Đề 5 — Ensemble Learning: Bagging cây quyết định"
        self.doc.core_properties.subject = "Tiểu luận giữa kỳ — Nhập môn học máy"
        self.doc.core_properties.keywords = "Bagging, bootstrap, hard voting, out-of-bag, Ensemble Learning"

    def page(self):
        self.doc.add_page_break()
        self.md.append("\n---\n")

    def heading(self, text, level=1):
        self.doc.add_heading(text, level=level)
        self.md += ["#" * (level + 1) + " " + text, ""]

    def p(self, text, bold=False, center=False):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        r.bold = bold
        self.md += [("**" + text + "**") if bold else text, ""]
        return p

    def bullet(self, text):
        p = self.doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(5)
        self.md += ["- " + text, ""]

    def equation(self, text):
        p = self.p(text, center=True)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(9)
        for run in p.runs:
            run.font.size = Pt(13)

    def table(self, headers, rows, widths=None):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        t.autofit = False
        for cell, value in zip(t.rows[0].cells, headers):
            cell.text = str(value)
            shade(cell, NAVY)
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
        header_repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(header_repeat)
        for i, values in enumerate(rows):
            row = t.add_row()
            for cell, value in zip(row.cells, values):
                cell.text = str(value)
                if i % 2 == 0:
                    shade(cell, "EEF5F3")
        for row in t.rows:
            cant_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant_split)
            for i, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if widths:
                    cell.width = Cm(widths[i])
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.size = Pt(10.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        def md_cell(value):
            return str(value).replace("|", r"\|").replace("\n", " / ")
        self.md.append("| " + " | ".join(md_cell(value) for value in headers) + " |")
        self.md.append("| " + " | ".join("---" for _ in headers) + " |")
        for values in rows:
            self.md.append("| " + " | ".join(md_cell(value) for value in values) + " |")
        self.md.append("")

    def code(self, content, language="python"):
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_together = True
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(.25)
        for i, line in enumerate(content.strip().splitlines()):
            r = p.add_run(("\n" if i else "") + line)
            r.font.name = "DejaVu Sans Mono"
            r.font.size = Pt(9)
        self.md += ["```" + language, content.strip(), "```", ""]

    def figure(self, name, caption, width=6.1):
        path = ROOT / "outputs" / name
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        p.paragraph_format.keep_with_next = True
        c = self.doc.add_paragraph(caption, style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.md += [f"![{caption}](../outputs/{name})", ""]

    def save(self):
        self.doc.save(REPORT / (STEM + ".docx"))
        (REPORT / (STEM + ".md")).write_text("\n".join(self.md), encoding="utf-8")


def normalized(text):
    return re.sub(r"\s+", " ", text).strip()


def contents_pages(pdf):
    if not pdf:
        return {}
    result = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True, text=True)
    pages = result.stdout.split("\f")
    found = {}
    # The contents itself repeats the headings, so skip the first two pages.
    for i, page in enumerate(pages[2:], start=3):
        page_text = normalized(page)
        for heading in TOC_HEADINGS:
            if heading not in found and normalized(heading) in page_text:
                found[heading] = i
    if set(found) != set(TOC_HEADINGS):
        raise SystemExit("Không tìm đủ tiêu đề trong PDF để cập nhật mục lục: " + str(set(TOC_HEADINGS) - set(found)))
    return found


def build(toc_pdf=None):
    report = json.loads((ROOT / "outputs/results.json").read_text(encoding="utf-8"))
    if (report.get("seed") != 42 or report.get("n_estimators") != 51
            or report.get("repeated_evaluation", {}).get("seeds") != list(range(42, 52))):
        raise SystemExit("Báo cáo cần kết quả đầy đủ mặc định: chạy de_5_ensemble_learning.py trước.")
    toy = build_toy_example()
    group = json.loads((REPORT / "group_info.json").read_text(encoding="utf-8"))
    members = group.get("members") or [{"name": "", "student_id": ""}]
    if len(members) > 3:
        raise SystemExit("Theo đề, nhóm tối đa 3 thành viên. Kiểm tra report/group_info.json.")
    scores = report["test_metrics"]
    single, bag, rf = scores
    summaries = report["repeated_evaluation"]["summary"]
    trace = report["disagreement_example"]
    toc = contents_pages(toc_pdf)
    p = Paper()
    p.md += ["# Đề 5 — Ensemble Learning", ""]
    # Cover: identity fields remain explicit until the group supplies them.
    p.doc.add_paragraph().paragraph_format.space_after = Pt(30)
    p.p("TIỂU LUẬN GIỮA KỲ", bold=True, center=True)
    p.p("HỌC PHẦN: NHẬP MÔN HỌC MÁY", bold=True, center=True)
    p.doc.add_paragraph().paragraph_format.space_after = Pt(28)
    title = p.p("ĐỀ 5 — ENSEMBLE LEARNING", bold=True, center=True)
    title.runs[0].font.size = Pt(24)
    title.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = p.p("Bagging cây quyết định:\ntừ ví dụ tính tay đến thực nghiệm phân loại", center=True)
    subtitle.runs[0].font.size = Pt(17)
    p.doc.add_paragraph().paragraph_format.space_after = Pt(28)
    if group.get("class_name"):
        p.p("Lớp: " + group["class_name"], center=True)
    if group.get("lecturer"):
        p.p("Giảng viên: " + group["lecturer"], center=True)
    p.table(["Thành viên", "Mã số sinh viên"], [
        [m.get("name") or "[ĐIỀN HỌ VÀ TÊN]", m.get("student_id") or "[ĐIỀN MSSV]"] for m in members
    ], [9, 7.5])
    p.p("Tháng 10 năm 2026", center=True)
    p.doc.core_properties.author = "; ".join(m["name"] for m in members if m.get("name")) or "Nhóm thực hiện — cần bổ sung thông tin"

    p.page()
    p.heading("MỤC LỤC")
    for heading in TOC_HEADINGS:
        line = heading + ("  ···  " + str(toc[heading]) if heading in toc else "")
        p.p(line)
    p.heading("Tóm tắt", level=2)
    p.p("Báo cáo tìm hiểu cách kết hợp nhiều mô hình để tạo một bộ phân loại. Trọng tâm là Bagging cây quyết định: lấy mẫu có hoàn lại, huấn luyện các cây và tổng hợp bằng đa số phiếu. Một ví dụ năm quan sát được tính xuyên suốt để thấy vì sao thay đổi mẫu huấn luyện có thể làm các cây đưa ra dự đoán khác nhau.")
    p.p(f"Phần thực nghiệm sử dụng dữ liệu Breast Cancer Wisconsin Diagnostic có sẵn trong scikit-learn. Bagging tự triển khai được đối chiếu với cây đơn và Random Forest trên 10 lần chia train/test có phân tầng. Accuracy trung bình của Bagging là {mean_sd(summaries[1], 'accuracy')}%, trong đó số sau dấu ± là độ lệch chuẩn mẫu tính theo điểm phần trăm. Kết quả cho thấy lợi ích trên dữ liệu khảo sát, đồng thời vẫn tồn tại lỗi và các giới hạn về đánh giá.")
    p.p("Từ khóa: Ensemble Learning; Bagging; bootstrap; hard voting; out-of-bag.")

    p.page()
    p.heading("MỞ ĐẦU")
    p.heading("Lý do chọn cách tiếp cận", level=2)
    p.p("Một cây quyết định dễ giải thích: mỗi nút đặt một câu hỏi về đặc trưng, quan sát đi theo nhánh phù hợp và nhận nhãn tại lá. Tuy nhiên, khi dữ liệu huấn luyện thay đổi, cấu trúc cây và dự đoán cũng có thể thay đổi đáng kể. Ensemble Learning khai thác nhiều mô hình thay vì phụ thuộc hoàn toàn vào một mô hình duy nhất [1, 2].")
    p.p("Bài chọn Bagging làm nội dung triển khai chính vì cơ chế của nó có thể quan sát trực tiếp: các mẫu bootstrap khác nhau tạo các cây khác nhau, sau đó đếm phiếu để ra quyết định. Cách làm này đủ đơn giản để tính tay nhưng vẫn thể hiện rõ bản chất của học tổ hợp.")
    p.heading("Mục tiêu và phạm vi", level=2)
    p.bullet("Giải thích vì sao cần nhiều mô hình, sự đa dạng giữa các mô hình và các cách tổng hợp dự đoán.")
    p.bullet("Trình bày trọn vẹn một ví dụ: dữ liệu → bootstrap → cây → dự đoán → đa số phiếu, kèm cách hiểu out-of-bag.")
    p.bullet("Tự triển khai bootstrap, hard voting và OOB; dùng cây quyết định của scikit-learn làm mô hình cơ sở.")
    p.bullet("Đánh giá trên dữ liệu có sẵn, so sánh với cây đơn và Random Forest, giải thích kết quả và lỗi thay vì chỉ báo cáo accuracy.")
    p.p("Boosting, Voting và Stacking được giới thiệu để định vị Bagging trong Ensemble Learning. Báo cáo không nhận đã triển khai hoặc đo thực nghiệm Boosting/Stacking. Cây quyết định cơ sở và Random Forest sử dụng thư viện; phần lấy mẫu và tổ hợp của ManualBagging được viết tường minh trong mã nguồn.")
    p.heading("Cách đọc báo cáo", level=2)
    p.p("Chương 1 trình bày nguyên lý. Chương 2 đi qua ví dụ tính tay. Chương 3 mô tả bài toán và cách đánh giá. Chương 4 giải thích mã nguồn, kết quả và nhận xét. Các công thức đều gắn với ý nghĩa hoặc con số cụ thể; phần thực nghiệm mang mục đích học thuật, không phải hệ thống chẩn đoán lâm sàng.")

    p.page()
    p.heading("1. NGUYÊN LÝ HOẠT ĐỘNG")
    p.heading("1.1. Ensemble Learning là gì?", level=2)
    p.p("Ensemble Learning kết hợp nhiều mô hình cơ sở thành một mô hình tổng hợp. Với phân loại, mỗi mô hình có thể đề xuất một nhãn hoặc một phân bố xác suất; bộ kết hợp dùng các đầu ra này để chọn nhãn cuối cùng. Ý tưởng không phải cứ có nhiều mô hình là tốt: lợi ích phụ thuộc cả chất lượng từng mô hình và việc chúng có mắc lỗi giống nhau hay không [2, 3].")
    p.p("Ví dụ, nếu ba mô hình cùng dự đoán sai ở đúng các quan sát giống nhau, bỏ phiếu không sửa được các lỗi đó. Nếu lỗi của chúng khác nhau và đa số thường đúng, tổ hợp có thể giảm tác động của một dự đoán sai riêng lẻ. Sự khác biệt hữu ích giữa các mô hình được gọi là tính đa dạng.")
    p.heading("1.2. Bagging: tạo đa dạng bằng dữ liệu", level=2)
    p.p("Bagging là viết tắt của bootstrap aggregating. Cho tập train có N quan sát, ta lấy N lượt có hoàn lại để tạo một tập bootstrap. Có hoàn lại nghĩa là một quan sát vừa được lấy vẫn có thể được lấy tiếp: một số quan sát bị lặp, một số quan sát không xuất hiện trong tập mới. Lặp cách lấy mẫu này B lần và huấn luyện một mô hình trên mỗi tập [1].")
    p.p("Các cây trong Bagging được học độc lập về mặt thuật toán. Một cây không cần chờ kết quả hay lỗi của cây trước để học. Do cây quyết định nhạy với thay đổi dữ liệu, các tập bootstrap có thể tạo những phép chia khác nhau. Sau khi học xong, cùng một quan sát mới được đưa qua tất cả cây rồi tổng hợp dự đoán.")
    p.p("Trực giác giảm phương sai: nếu một thay đổi nhỏ của train làm một cây thay đổi mạnh, tổng hợp nhiều cây khác nhau có thể ổn định hơn. Bagging thường phù hợp với loại mô hình có tính biến động cao như cây quyết định. Báo cáo không suy từ trực giác này rằng mọi tổ hợp đều tốt hơn mô hình đơn.")
    p.heading("1.3. Quy tắc hard voting", level=2)
    p.equation("H(x) = argmax theo lớp c của tổng I[h_b(x) = c], b = 1,…,B")
    p.p("Trong đó h_b(x) là nhãn do cây b dự đoán; I bằng 1 nếu cây chọn lớp c và bằng 0 nếu không. Tổng các I chính là số phiếu của lớp c; argmax chọn lớp có nhiều phiếu nhất. Với hai phiếu ác tính và một phiếu lành tính, tổ hợp chọn ác tính. Khi hòa phiếu, mã chọn lớp đầu trong classes_; 51 cây tránh hòa ở bài toán hai lớp, nhưng số cây lẻ vẫn có thể hòa khi có nhiều hơn hai lớp.")

    p.page()
    p.heading("1.4. Hard voting và soft voting", level=2)
    p.p("Hard voting chỉ đếm nhãn. Soft voting lấy trung bình xác suất dự đoán của các mô hình rồi chọn lớp có xác suất trung bình lớn nhất. Hai cách có thể cho kết quả khác nhau; soft voting cũng không tự động tốt hơn nếu các xác suất đầu vào không đáng tin cậy [3].")
    p.table(["Mô hình", "P(ác tính | x)", "Nhãn riêng"], [["1", "0,51", "Ác tính"], ["2", "0,51", "Ác tính"], ["3", "0,01", "Lành tính"]], [4, 6, 6.5])
    p.p("Trong ví dụ giả định này, hard voting chọn ác tính vì có 2/3 phiếu. Soft voting tính (0,51 + 0,51 + 0,01)/3 ≈ 0,343 cho ác tính và 0,657 cho lành tính, nên chọn lành tính. Các số này minh họa quy tắc tổng hợp, không phải đầu ra của thực nghiệm Breast Cancer.")
    p.heading("1.5. Những cách kết hợp thường gặp", level=2)
    p.table(["Cách tiếp cận", "Ý tưởng chính", "Vai trò trong bài"], [
        ["Bagging", "Học độc lập trên các bootstrap rồi tổng hợp", "Trọng tâm tự triển khai"],
        ["Boosting", "Học tuần tự; mô hình sau xử lý sai sót hoặc giảm loss", "Giới thiệu để đối chiếu"],
        ["Voting", "Kết hợp nhãn hoặc xác suất của nhiều mô hình", "Hard voting của Bagging"],
        ["Stacking", "Dùng dự đoán ngoài fold để học bộ kết hợp", "Giới thiệu để đối chiếu"],
    ], [3, 9, 4.5])
    p.p("Các cách trên không hoàn toàn tách rời: Bagging cũng cần quy tắc tổng hợp như voting. Với AdaBoost, một ví dụ của Boosting, các mẫu đang bị dự đoán sai nhận trọng số lớn hơn để mô hình sau chú ý hơn; trọng số được chuẩn hóa sau mỗi vòng [6]. Với Stacking, tầng kết hợp phải được học từ dự đoán ngoài fold của các mô hình cơ sở, tránh dùng dự đoán quá tốt ngay trên dữ liệu chúng đã học [7].")
    p.heading("1.6. Random Forest và giới hạn của so sánh", level=2)
    p.p("Random Forest kết hợp bootstrap với việc chọn ngẫu nhiên một tập con đặc trưng tại từng phép chia, giúp tăng đa dạng giữa các cây [2]. Trong scikit-learn, RandomForestClassifier tổng hợp bằng trung bình xác suất lớp của các cây, còn ManualBagging trong bài tổng hợp bằng đa số nhãn. Do đó, đối chiếu hai mô hình không phải thí nghiệm chỉ thay đổi riêng một yếu tố chọn đặc trưng [3].")

    p.page()
    p.heading("2. QUY TRÌNH VÀ VÍ DỤ TÍNH TAY")
    p.heading("2.1. Quy trình Bagging", level=2)
    p.table(["Bước", "Thao tác", "Điều cần giữ đúng"], [
        ["1", "Chuẩn bị train và giữ riêng test", "Test không tham gia học hay chọn cấu hình"],
        ["2", "Lấy B bootstrap, mỗi tập N lượt", "Lấy có hoàn lại, không chia train thành B phần"],
        ["3", "Huấn luyện một cây trên mỗi bootstrap", "Các cây học độc lập"],
        ["4", "Đưa cùng x qua mọi cây", "Tất cả cây dự đoán cùng một quan sát"],
        ["5", "Đếm phiếu và chọn lớp đa số", "Quy tắc hòa phiếu phải rõ ràng"],
        ["6", "Đánh giá test và đọc các loại lỗi", "Không chỉ nhìn accuracy"],
    ], [1.4, 7.7, 7.4])
    p.heading("2.2. Dữ liệu minh họa xuyên suốt", level=2)
    p.p("Dùng năm quan sát một chiều để có thể theo dõi từng con số. Đây là dữ liệu giả lập; nhãn ác/lành chỉ giúp thống nhất cách gọi, không mô tả đặc trưng y tế thực tế. Chỉ số hàng bắt đầu từ 0. Cây trong ví dụ chỉ được chia một lần (max_depth=1).")
    p.table(["Chỉ số", "0", "1", "2", "3", "4"], [["Đặc trưng x", "1", "2", "3", "4", "5"], ["Nhãn", "Lành (1)", "Lành (1)", "Ác (0)", "Ác (0)", "Ác (0)"]], [3, 2.7, 2.7, 2.7, 2.7, 2.7])
    p.p("Ta cố định ba bootstrap dưới đây để người đọc tái lập đúng phép tính. Trong thực nghiệm thật, chương trình sinh ngẫu nhiên các bootstrap theo seed; các chỉ số được cố định ở đây chỉ phục vụ giảng giải.")
    p.table(["Cây", "Chỉ số bootstrap", "Hàng không có mặt"], [[t["name"], str(t["bootstrap_indices"]), str(t["omitted_indices"])] for t in toy["trees"]], [2, 9, 5.5])

    p.page()
    p.heading("2.3. Từ bootstrap đến phép chia của cây", level=2)
    p.p("Cây quyết định chọn phép chia làm các nhánh thuần nhất hơn. Với hai lớp, độ hỗn tạp Gini tại một nút được tính bằng 1 − p_ác² − p_lành². Nút chỉ có một lớp có Gini bằng 0. Gini sau chia là trung bình có trọng số theo số quan sát của hai nhánh [5].")
    p.equation("Gini sau chia = (N_trái/N) × Gini_trái + (N_phải/N) × Gini_phải")
    p.p("Cây h1 dùng [0,0,2,4,1], tức các giá trị x là [1,1,3,5,2], với ba nhãn lành và hai nhãn ác. Gini trước chia bằng 1 − (3/5)² − (2/5)² = 0,48. Nếu chọn ngưỡng x ≤ 2,5, nhánh trái có ba mẫu lành; nhánh phải có hai mẫu ác. Gini mỗi nhánh bằng 0, nên Gini sau chia bằng (3/5)×0 + (2/5)×0 = 0. Mức giảm bằng 0,48.")
    p.p("Vì sao ngưỡng là 2,5? Sau khi sắp các giá trị khác nhau thành 1, 2, 3, 5, các ngưỡng ứng viên nằm giữa hai giá trị liên tiếp. Ngưỡng 2,5 tách hoàn toàn hai lớp nên đạt mức giảm lớn nhất ở nút này. Các ngưỡng 1,5 và 4 không cho hai nhánh đều thuần nhất.")
    p.table(["Cây", "Nhãn trong bootstrap", "Ngưỡng được học", "Gini trước → sau"], [
        [t["name"], f"{t['left_size']} lành; {t['right_size']} ác", f"x ≤ {number(t['threshold'], 1)}", f"{number(t['gini_before'])} → {number(t['gini_after'])}"]
        for t in toy["trees"]
    ], [1.5, 5, 4.5, 5.5])
    p.p("Đối với h2, bootstrap không có hàng x=3; hai giá trị sát ranh giới lớp là x=2 và x=4, nên cây chọn trung điểm 3,0. Cây h3 vẫn có x=2 và x=3, nên ngưỡng là 2,5. Cả ba cây đều phân loại đúng các mẫu trong bootstrap của riêng mình, nhưng ranh giới chúng học không hoàn toàn giống nhau. Đây là cơ chế tạo đa dạng mà ví dụ cần làm rõ.")
    p.p("Các ngưỡng và giá trị Gini trong bảng được kiểm tra bằng DecisionTreeClassifier(max_depth=1). Ví dụ không giả định trước phiếu của cây: phiếu ở bước tiếp theo được lấy từ chính ba cây vừa huấn luyện.")

    p.page()
    p.heading("2.4. Dự đoán cùng một quan sát và cộng phiếu", level=2)
    p.p("Đưa cùng quan sát mới x=2,8 qua ba cây. Quy tắc mỗi cây là: x không vượt ngưỡng thì chọn lành; x lớn hơn ngưỡng thì chọn ác. Chưa biết nhãn thật của quan sát mới này, nên chỉ kết luận dự đoán, không tự gọi dự đoán là đúng.")
    p.table(["Cây", "So sánh", "Dự đoán"], [
        [t["name"], f"2,8 {'>' if t['query_prediction'] == 0 else '≤'} {number(t['threshold'], 1)}", LABEL[t["query_prediction"]]]
        for t in toy["trees"]
    ], [2, 7, 7.5])
    p.equation("Số phiếu ác = 2; số phiếu lành = 1 → H(2,8) = ác tính (0)")
    p.p("Như vậy, dữ liệu và bootstrap đã dẫn tới ba ngưỡng thực, ba dự đoán thực rồi mới đến kết quả voting. Tỷ lệ 2/3 chỉ là tỷ lệ cây bỏ phiếu ác trong ví dụ; nó không phải xác suất mắc bệnh đã được hiệu chỉnh.")
    p.heading("2.5. Out-of-bag: dùng mẫu cây chưa học", level=2)
    p.p("Một quan sát không có trong bootstrap của cây được gọi là out-of-bag (OOB) đối với cây đó. Khi đánh giá OOB cho một hàng train, chỉ những cây không học hàng ấy mới được bỏ phiếu. Không được dùng mọi cây để đánh giá chính dữ liệu train rồi gọi đó là OOB.")
    p.table(["Hàng", "Cây đủ điều kiện OOB", "Thật → dự đoán", "Kết quả"], [
        ["0 (x=1)", "h3", "Lành → Lành", "Đúng"],
        ["2 (x=3)", "h2", "Ác → Lành", "Sai"],
        ["3 (x=4)", "h1", "Ác → Ác", "Đúng"],
        ["1 và 4", "Không có cây OOB", "Không dự đoán", "Loại khỏi accuracy OOB"],
    ], [2.6, 5.2, 4.4, 4.3])
    p.p("Ở ví dụ nhỏ, coverage OOB = 3/5 = 60%; accuracy OOB trên ba hàng có phiếu = 2/3 ≈ 66,67%. Hàng 1 và 4 không được gán một lớp giả để tính điểm. Với N quan sát, xác suất một hàng không xuất hiện sau N lượt lấy mẫu là (1−1/N)^N; khi N lớn tiến gần 36,8%. Với N=5, giá trị chính xác là 32,768% [1].")

    p.page()
    p.heading("3. BÀI TOÁN ỨNG DỤNG VÀ THIẾT KẾ THỰC NGHIỆM")
    p.heading("3.1. Mục tiêu, dữ liệu đầu vào và đầu ra", level=2)
    p.p("Bài toán minh họa là phân loại mẫu khối u vú thành ác tính hoặc lành tính từ các đặc trưng đã trích xuất. Dữ liệu Breast Cancer Wisconsin Diagnostic được nạp bằng load_breast_cancer của scikit-learn [4]. Ta học ánh xạ từ một vector 30 giá trị số sang một trong hai nhãn.")
    p.table(["Thuộc tính", "Mô tả"], [
        ["Số quan sát", "569"], ["Đầu vào", "30 đặc trưng số của nhân tế bào, như radius, texture, perimeter, area"],
        ["Đầu ra", "0 = malignant (ác tính); 1 = benign (lành tính)"],
        ["Số lượng lớp", "212 ác tính; 357 lành tính"],
        ["Nguồn dùng trong mã", "sklearn.datasets.load_breast_cancer; không cần tải dữ liệu riêng"],
        ["Mục đích", "Minh họa phương pháp học máy, chưa xác thực lâm sàng"],
    ], [4.5, 12])
    p.p("Đầu vào không phải ảnh gốc và bài không huấn luyện mô hình nhận dạng ảnh. Các đặc trưng đã được chuẩn bị; cây sử dụng phép chia theo ngưỡng nên không yêu cầu chuẩn hóa thang đo như các mô hình dựa trên khoảng cách. Đây là lựa chọn phù hợp với một bài minh họa tập trung vào cơ chế tổ hợp.")
    p.heading("3.2. Chia dữ liệu và giữ test độc lập với lựa chọn", level=2)
    p.p(f"Mỗi seed chia 75% train và 25% test, có stratify theo nhãn để giữ gần tỷ lệ hai lớp. Tại seed {report['seed']}, train có {report['train_size']} mẫu (159 ác và 267 lành), test có {report['test_size']} mẫu (53 ác và 90 lành). Test không được dùng để chọn độ sâu, số cây hay seed tốt nhất.")
    p.p("Trên train, thử max_depth trong [2,3,4,5,6] bằng StratifiedKFold năm fold. Mỗi lần thử học trên bốn fold và chấm trên fold còn lại, rồi lấy trung bình balanced accuracy. Cấu hình có điểm cao nhất được chọn; nếu hòa, mã ưu tiên độ sâu nhỏ hơn. Sau đó huấn luyện mô hình trên toàn train và mới đánh giá test.")

    p.page()
    p.heading("3.3. Cấu hình đối chiếu và lặp thí nghiệm", level=2)
    p.table(["Mô hình", "Cấu hình", "Cách tổng hợp"], [
        ["Cây đơn", "Một DecisionTreeClassifier", "Dự đoán trực tiếp"],
        ["Manual Bagging", "51 cây; bootstrap tự viết", "Đa số nhãn"],
        ["Random Forest", "51 cây; max_features='sqrt'", "Trung bình xác suất cây"],
    ], [3.6, 7.2, 5.7])
    p.p("Ba mô hình dùng chung train/test và độ sâu vừa chọn bằng cây đơn để đối chiếu đơn giản. Đây chưa phải cấu hình tối ưu riêng cho từng ensemble. Với 30 đặc trưng, sqrt tương ứng mục tiêu 5 đặc trưng tại một phép chia của Random Forest; thư viện có thể xem thêm đặc trưng để tìm phép chia hợp lệ [3].")
    p.p("Thí nghiệm lặp cho các seed cố định 42–51. Mỗi seed chia lại dữ liệu và chọn lại độ sâu trên train của chính seed đó. Báo cáo giữ seed 42 làm lần tham chiếu để phân tích lỗi cụ thể; bảng tổng hợp sử dụng đủ cả mười seed, không chọn lượt có điểm đẹp nhất.")
    p.heading("3.4. Chỉ số và quy ước lớp dương", level=2)
    p.p("Chọn ác tính = 0 làm lớp dương. Vì vậy, không được dùng mặc định pos_label=1 rồi gọi đó là recall ác tính. TP là mẫu ác được dự đoán ác; FN là mẫu ác bị dự đoán lành; FP là mẫu lành bị dự đoán ác; TN là mẫu lành được dự đoán lành. Ma trận có hàng là nhãn thật, cột là dự đoán, thứ tự [ác, lành].")
    p.table(["Chỉ số", "Công thức", "Cách đọc"], [
        ["Accuracy", "(TP+TN)/N", "Tỷ lệ dự đoán đúng tổng thể"],
        ["Precision ác", "TP/(TP+FP)", "Trong các dự đoán ác, bao nhiêu là ác thật"],
        ["Recall ác", "TP/(TP+FN)", "Trong các mẫu ác thật, phát hiện được bao nhiêu"],
        ["F1 ác", "2TP/(2TP+FP+FN)", "Cân bằng precision và recall"],
        ["Balanced accuracy", "(Recall ác + Recall lành)/2", "Cân bằng đóng góp hai lớp; dùng chọn depth"],
    ], [3.6, 5.2, 7.7])
    p.p("Recall ác tính và số FN được chú ý vì chúng thể hiện các ca ác tính bị bỏ sót. Điều này không biến mô hình thành công cụ y tế đã được kiểm định. Khi tổng hợp nhiều seed, báo cáo mean và SD mẫu (ddof=1). Các tập test giữa các seed có thể giao nhau, nên SD chỉ mô tả dao động, không phải khoảng tin cậy độc lập.")

    p.page()
    p.heading("4. TRIỂN KHAI MÃ NGUỒN VÀ KẾT QUẢ")
    p.heading("4.1. Tổ chức mã và phần tự triển khai", level=2)
    p.table(["Tệp / thành phần", "Nhiệm vụ"], [
        ["ensemble_model.py / ManualBagging", "Kiểm tra dữ liệu; lấy mẫu bootstrap; học cây; đếm phiếu; OOB"],
        ["de_5_ensemble_learning.py", "Chia train/test; chọn depth; so sánh ba mô hình; xuất kết quả và biểu đồ"],
        ["examples/toy_bagging.py", "Tái lập ví dụ năm quan sát và các phép tính trong Chương 2"],
        ["tests/", "Kiểm thử voting/OOB, tái lập, quy trình và xuất tệp"],
    ], [7, 9.5])
    p.p("ManualBagging kế thừa giao diện estimator của scikit-learn để dùng được clone và cross-validation, nhưng không gọi BaggingClassifier để thực hiện thay việc lấy mẫu hay tổng hợp. Nhãn đầu vào được ánh xạ về các chỉ số liên tiếp bằng np.unique, giúp đếm phiếu đúng cả khi nhãn là chuỗi hoặc số không liên tiếp.")
    p.heading("4.2. Lấy mẫu và huấn luyện", level=2)
    p.code("""for _ in range(self.n_estimators):
    indices = rng.integers(0, n_samples, size=n_samples)
    tree = DecisionTreeClassifier(max_depth=self.max_depth,
                                  random_state=tree_seed)
    tree.fit(X[indices], encoded[indices])
    self.estimators_.append(tree)""")
    p.p("Đây là trích lược để giải thích; tree_seed trong bản đầy đủ được sinh từ rng. rng.integers tạo đúng N chỉ số có hoàn lại. X và y phải lấy cùng indices để không lệch nhãn. Mỗi cây có seed riêng để tái lập; bản đầy đủ còn lưu các chỉ số bootstrap nhằm kiểm tra và tính OOB. Cây cơ sở do thư viện cung cấp, phần vòng lặp Bagging thuộc mã tự triển khai.")
    p.heading("4.3. Đếm phiếu và chọn nhãn", level=2)
    p.code("""counts = np.zeros((len(X), len(self.classes_)), dtype=np.int64)
rows = np.arange(len(X))
for tree in self.estimators_:
    predicted = tree.predict(X).astype(int)
    np.add.at(counts, (rows, predicted), 1)
return self.classes_[counts.argmax(axis=1)]""")
    p.p("Mỗi hàng của counts ứng với một quan sát, mỗi cột ứng với một lớp. Mỗi cây tăng đúng một phiếu cho mỗi quan sát. argmax tìm cột có phiếu lớn nhất, sau đó classes_ ánh xạ về nhãn gốc. predict_vote_fraction chia các số phiếu cho số cây; tên hàm nhấn mạnh đây là tỷ lệ phiếu, không nhận là xác suất đã hiệu chỉnh.")

    p.page()
    p.heading("4.4. OOB và các kiểm tra giúp kết quả đáng tin", level=2)
    p.p("Với mỗi bootstrap, np.bincount đếm số lần mỗi hàng xuất hiện. Các hàng có số đếm bằng 0 mới được cây dự đoán để tích lũy phiếu OOB. Khi cộng xong mọi cây, chỉ những hàng có ít nhất một phiếu mới tham gia accuracy OOB. Hàng chưa có phiếu giữ giá trị NaN trong oob_decision_function_; không coi chúng là dự đoán lớp 0.")
    p.p("Độ sâu đã được chọn bằng CV trên toàn train trước khi tính OOB. Vì thế OOB sau bước lựa chọn này không hoàn toàn độc lập với tuning; đây là chẩn đoán bổ sung. Kết quả test được giữ riêng vẫn là đánh giá cuối của mỗi seed.")
    p.table(["Nhóm kiểm tra", "Hành vi được kiểm tra"], [
        ["Dữ liệu và giao diện", "Đầu vào sai; dự đoán trước fit; ánh xạ nhãn; clone và CV"],
        ["Bootstrap và voting", "Lấy có hoàn lại; mỗi cây đúng N lượt; tổng phiếu; xử lý hòa"],
        ["OOB", "Chỉ cây bỏ qua mẫu mới có phiếu; coverage; hàng chưa có phiếu"],
        ["Quy trình thí nghiệm", "Chọn depth chỉ nhận train; hồi quy kết quả seed 42"],
        ["Xuất kết quả", "CLI chạy thực; JSON chuẩn; CSV; ảnh PNG; số cây tùy chỉnh"],
        ["Ví dụ tính tay", "Ngưỡng/Gini thực; cùng x cho mọi cây; đúng nguồn phiếu OOB"],
    ], [4.5, 12])
    p.heading("4.5. Cách tái lập", level=2)
    p.p("Dùng Python 3.12 và cài requirements.txt. Các phiên bản trực tiếp của lần chạy được ghi trong results.json; việc ghim phiên bản cùng seed giúp hạn chế sai khác giữa môi trường. Dữ liệu đã nằm trong scikit-learn, không cần API key.")
    p.code("""python3 -m venv .venv
.venv/bin/python -m pip install -r requirements.txt
.venv/bin/python -m unittest discover -s tests -v
.venv/bin/python examples/toy_bagging.py --output-dir outputs/toy
.venv/bin/python de_5_ensemble_learning.py""", language="bash")
    p.p("Trên Windows, dùng .venv\\Scripts\\python.exe thay cho .venv/bin/python. Trong cloud có home chỉ đọc, đặt MPLCONFIGDIR=/tmp/hocmaygk-matplotlib và XDG_CACHE_HOME=/tmp/hocmaygk-cache trước khi chạy. Demo nhanh dùng --repeats 1 --output-dir outputs/quick để không ghi đè kết quả đầy đủ. Bảng/biểu đồ trong báo cáo này được lấy từ cấu hình mặc định 10 seed, không suy ra từ demo một seed.")

    p.page()
    p.heading("4.6. Kết quả lần tham chiếu: seed 42", level=2)
    p.p(f"Train có {report['train_size']} mẫu, test có {report['test_size']} mẫu. Độ sâu {report['selected_depth']} được chọn vì balanced accuracy CV trung bình đạt {pct(report['depth_selection_train_cv_balanced_accuracy']['3'])}, cao nhất trong năm độ sâu ứng viên. Hai ensemble đều có {report['n_estimators']} cây.")
    p.table(["Mô hình", "Accuracy", "Precision ác", "Recall ác", "F1 ác"], [[MODEL[r["model"]], pct(r["accuracy"]), pct(r["precision_malignant"]), pct(r["recall_malignant"]), pct(r["f1_malignant"])] for r in scores], [5.3, 2.8, 2.8, 2.8, 2.8])
    p.figure("confusion_matrices.png", "Hình 1. Ma trận nhầm lẫn của ba mô hình, seed 42; hàng thật, cột dự đoán.", width=6.35)
    p.p(f"Bagging dự đoán đúng {bag['tp'] + bag['tn']}/{report['test_size']} mẫu, cây đơn đúng {single['tp'] + single['tn']}/{report['test_size']}. Chênh lệch là một mẫu, tương đương {number((bag['accuracy'] - single['accuracy']) * 100)} điểm phần trăm accuracy. Không nên mô tả một chênh lệch nhỏ trên một lần chia như bằng chứng ưu thế tổng quát.")
    p.heading("Thay số để kiểm tra chỉ số Bagging", level=3)
    p.table(["Chỉ số", "Thay số", "Kết quả"], [
        ["Accuracy", "(49+87)/143", pct(bag["accuracy"])],
        ["Precision ác", "49/(49+3)", pct(bag["precision_malignant"])],
        ["Recall ác", "49/(49+4)", pct(bag["recall_malignant"])],
        ["F1 ác", "2×49/(2×49+3+4)", pct(bag["f1_malignant"])],
    ], [4.7, 7.1, 4.7])
    p.p("Cây đơn bỏ sót 5/53 mẫu ác; Bagging còn bỏ sót 4/53 mẫu ác. Random Forest trùng các chỉ số tổng hợp của Bagging trong lượt này. Trùng số đếm không đủ để kết luận hai phương pháp tương đương hoặc dự đoán giống nhau trên mọi quan sát.")

    p.page()
    p.heading("4.7. Độ ổn định trên mười lần chia", level=2)
    p.p("Bảng dưới lấy đủ các seed 42–51. Mỗi số trước dấu ± là trung bình theo %, số sau là SD mẫu theo điểm phần trăm. Vì các test có thể giao nhau, không diễn giải bảng thành khoảng tin cậy hay một kiểm định ý nghĩa thống kê độc lập.")
    p.table(["Mô hình", "Accuracy (%)", "Recall ác (%)", "F1 ác (%)"], [[MODEL[r["model"]], mean_sd(r, "accuracy"), mean_sd(r, "recall_malignant"), mean_sd(r, "f1_malignant")] for r in summaries], [5.3, 3.7, 3.7, 3.8])
    p.figure("stability.png", "Hình 2. Trung bình accuracy và recall qua 10 holdout; thanh sai số là SD mẫu.", width=6.25)
    p.p("Bagging có accuracy trung bình cao hơn cây đơn khoảng 1,96 điểm phần trăm và SD accuracy thấp hơn trong khảo sát này. Random Forest có accuracy trung bình cao hơn nhẹ so với Bagging, trong khi recall ác tính trung bình của Bagging cao hơn nhẹ. Vì vậy, nhận xét cần gắn với chỉ số đang xem, không chỉ tuyên bố một mô hình thắng mọi mặt.")
    p.p("Sự cải thiện có thể được giải thích bởi việc tổng hợp các cây học trên các mẫu khác nhau làm giảm tác động của một cây sai riêng lẻ. Tuy nhiên, thực nghiệm này không đo trực tiếp tương quan lỗi hay tách riêng tác dụng của từng thành phần. Đó là cách giải thích phù hợp với cơ chế, chưa phải kết luận nhân quả đã được kiểm chứng riêng.")
    p.heading("OOB của Bagging ở seed tham chiếu", level=3)
    p.p(f"OOB accuracy = 401/426 = {pct(report['oob']['accuracy'])}; coverage = {pct(report['oob']['coverage'])}. Mỗi hàng nhận từ {report['oob']['min_trees_per_sample']} đến {report['oob']['max_trees_per_sample']} phiếu từ những cây chưa thấy nó. Số phiếu không giống nhau giữa các hàng và không phải đủ 51 phiếu. OOB được đọc cùng coverage và giới hạn liên quan đến bước chọn depth đã nêu.")

    p.page()
    p.heading("4.8. Một trường hợp tổ hợp vẫn sai", level=2)
    p.p(f"Để nhìn rõ cơ chế thay vì chọn ví dụ đẹp, chương trình lấy mẫu test đầu tiên có phiếu ở cả hai lớp. Mẫu tại vị trí test {trace['test_position']} ứng với hàng {trace['dataset_row']} trong dữ liệu gốc, đều đếm từ 0; đây là chỉ số hàng, không phải mã bệnh nhân. Nhãn thật là lành tính.")
    p.figure("voting_example.png", "Hình 3. Ví dụ bất đồng thật: 27 phiếu ác, 24 phiếu lành; tổ hợp dự đoán sai.", width=5.5)
    p.p("Có 27 cây chọn ác tính và 24 cây chọn lành tính, nên hard voting chọn ác tính. Do nhãn thật là lành tính, đây là dương tính giả (FP). Ví dụ chứng minh việc có đa số phiếu không đảm bảo dự đoán đúng. Tỷ lệ 27/51 không được diễn giải là xác suất mắc bệnh đã hiệu chỉnh.")
    p.heading("4.9. Giới hạn và hướng cải thiện vừa sức", level=2)
    p.bullet("Độ sâu được chọn bằng cây đơn rồi dùng chung. Có thể cải thiện bằng CV riêng cho mỗi mô hình, vẫn giữ test ngoài quá trình chọn cấu hình.")
    p.bullet("51 cây là cấu hình minh họa cố định, chưa chứng minh tối ưu. Khảo sát số cây bằng train/CV là hướng mở rộng hợp lý; không chọn dựa trên test.")
    p.bullet("Mười holdout có giao nhau; SD chỉ mô tả dao động. Cần dữ liệu ngoài hoặc thiết kế đánh giá khác để kết luận rộng hơn.")
    p.bullet("Dữ liệu hiện có chưa đại diện mọi quần thể thực tế. Kết quả học thuật và các ca bỏ sót còn tồn tại không đủ để sử dụng cho quyết định lâm sàng.")

    p.page()
    p.heading("KẾT LUẬN")
    p.p("Báo cáo đã trình bày Ensemble Learning qua ý tưởng kết hợp những mô hình vừa đủ tốt vừa có sự đa dạng. Bagging hiện thực ý tưởng đó bằng bootstrap và tổng hợp dự đoán. Ví dụ năm quan sát nối trực tiếp dữ liệu, các tập bootstrap, ngưỡng cây, phiếu của cùng một mẫu và kết quả cuối cùng; phép tính OOB còn cho thấy vì sao phải phân biệt mẫu có và không có phiếu hợp lệ.")
    p.p("Mã ManualBagging tự thực hiện lấy mẫu, hard voting và OOB, trong khi sử dụng cây quyết định của scikit-learn làm thành phần cơ sở. Quy trình thực nghiệm giữ test ngoài việc chọn độ sâu, so sánh trên cùng các lần chia và xuất dữ liệu đủ để kiểm tra kết quả. Cách triển khai vì vậy thể hiện được cơ chế tổ hợp thay vì chỉ gọi sẵn một mô hình Bagging.")
    p.p(f"Trên mười holdout đã khảo sát, accuracy trung bình của Bagging là {mean_sd(summaries[1], 'accuracy')}%, cao hơn cây đơn trong bộ thí nghiệm này. Kết quả seed 42 cho thấy Bagging vẫn bỏ sót bốn ca ác tính và ví dụ bất đồng cho thấy tổ hợp vẫn có thể sai. Kết luận thích hợp là Bagging có lợi ích trên dữ liệu và cấu hình đã khảo sát, đồng thời cần giữ giới hạn khi suy rộng.")
    p.heading("Ba ý chính để trình bày", level=2)
    p.bullet("Bootstrap tạo các tập học khác nhau; đó là nguồn đa dạng của các cây trong Bagging.")
    p.bullet("Hard voting đếm phiếu của cùng một quan sát; đa số phiếu không đồng nghĩa chắc chắn đúng hay xác suất bệnh.")
    p.bullet("Đánh giá phải tách train/test, đọc các loại lỗi và mô tả trung thực phạm vi của kết quả.")
    p.heading("Dữ liệu kiểm chứng đi kèm", level=2)
    p.p("results.json lưu cấu hình, phiên bản, điểm CV, OOB và trace; metrics.csv lưu lần tham chiếu; repeated_metrics.csv và summary_metrics.csv lưu điểm từng seed và thống kê; predictions.csv lưu nhãn thật, dự đoán và số phiếu từng mẫu test. Các ảnh trong báo cáo được sinh từ cùng lượt chạy đầy đủ. examples/toy_bagging.py tái lập phép tính của Chương 2.")

    p.page()
    p.heading("TÀI LIỆU THAM KHẢO")
    refs = [
        "[1] Breiman, L. (1996). Bagging Predictors. Machine Learning, 24, 123–140. https://doi.org/10.1007/BF00058655",
        "[2] Breiman, L. (2001). Random Forests. Machine Learning, 45, 5–32. https://doi.org/10.1023/A:1010933404324",
        "[3] scikit-learn developers. Ensemble methods. Tài liệu hướng dẫn chính thức. https://scikit-learn.org/stable/modules/ensemble.html",
        "[4] scikit-learn developers. load_breast_cancer — Breast Cancer Wisconsin (Diagnostic). Tài liệu API và mô tả dữ liệu. https://scikit-learn.org/stable/modules/generated/sklearn.datasets.load_breast_cancer.html",
        "[5] scikit-learn developers. Decision Trees. Tài liệu về tiêu chí chia, Gini và triển khai cây. https://scikit-learn.org/stable/modules/tree.html",
        "[6] Freund, Y., & Schapire, R. E. (1997). A Decision-Theoretic Generalization of On-Line Learning and an Application to Boosting. Journal of Computer and System Sciences, 55(1), 119–139. https://doi.org/10.1006/jcss.1997.1504",
        "[7] Wolpert, D. H. (1992). Stacked Generalization. Neural Networks, 5(2), 241–259. https://doi.org/10.1016/S0893-6080(05)80023-1",
    ]
    for ref in refs:
        p.p(ref)
    p.heading("Nguồn của kết quả và ranh giới đóng góp", level=2)
    p.p("Các bảng số liệu và hình thực nghiệm do mã trong repository sinh từ dữ liệu của scikit-learn, không lấy điểm benchmark từ các bài báo trên. Nguồn [1] và [2] cung cấp cơ sở phương pháp; [3]–[5] giải thích hành vi thư viện và dữ liệu; [6]–[7] phục vụ phần tổng quan. Cây cơ sở và Random Forest thuộc scikit-learn; mã minh họa tập trung vào bootstrap, tổng hợp hard-vote và kiểm tra OOB.")
    p.p("Thông tin phiên bản của lần chạy: Python " + report["versions"]["python"] + "; NumPy " + report["versions"]["numpy"] + "; scikit-learn " + report["versions"]["scikit-learn"] + "; matplotlib " + report["versions"]["matplotlib"] + ".")
    p.save()
    print("Saved", REPORT / (STEM + ".docx"))
    print("Saved", REPORT / (STEM + ".md"))
    if toc:
        print("TOC page numbers read from PDF:", toc)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--toc-pdf", type=Path, help="PDF xuất lần trước để cập nhật số trang mục lục")
    build(parser.parse_args().toc_pdf)
